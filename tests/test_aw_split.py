"""
Verification for the AW Consultancy agency/agent commission split.

Numbers cross-checked against the real sample file: PO 20260266
(Agency AC108-02, Net Price RM18,300) shows a real 1st Half Commission
split of RM640.50 (agency, 3.5%) and RM732.00 (agent, 4%) - used below
as the ground truth, not an invented example.

Run with: pytest tests/test_aw_split.py -v
"""

import datetime

import openpyxl

from app.commission import calculate_agency_agent_split
from app.pipeline import process_upload, generate_report
from app.db.connection import get_connection
from tests.helpers import build_master_report, confirm_all_pending
from app import rules


def _db_path(tmp_path):
    return str(tmp_path / "ledger.db")


def test_split_matches_the_real_confirmed_numbers():
    agency_amount, agent_amount, total = calculate_agency_agent_split(
        net_price=18300,
        agency_pct=rules.AW_AGENCY_INSTALLMENT_PCT,
        agent_pct=rules.AW_AGENT_INSTALLMENT_PCT,
        fb_lead_referred=False,
        deduction_pct=rules.AW_FB_LEAD_DEDUCTION_INSTALLMENT_PCT,
    )
    assert agency_amount == 640.50
    assert agent_amount == 732.00
    assert total == 1372.50


def test_fb_lead_deduction_reduces_only_the_agency_share():
    without_deduction = calculate_agency_agent_split(
        net_price=18300, agency_pct=0.035, agent_pct=0.04,
        fb_lead_referred=False, deduction_pct=0.015,
    )
    with_deduction = calculate_agency_agent_split(
        net_price=18300, agency_pct=0.035, agent_pct=0.04,
        fb_lead_referred=True, deduction_pct=0.015,
    )
    # Agent share is identical either way.
    assert without_deduction[1] == with_deduction[1] == 732.00
    # Agency share drops by exactly the deduction (1.5% of 18300 = 274.50).
    assert without_deduction[0] == 640.50
    assert with_deduction[0] == 366.00
    assert with_deduction[0] == round(without_deduction[0] - 274.50, 2)
    # Total reflects the reduced agency share.
    assert with_deduction[2] == round(with_deduction[0] + with_deduction[1], 2)


def test_full_payment_split_via_full_pipeline(tmp_path):
    """
    End-to-end: uploading a Master report with an AW Consultancy agency
    code produces a commission_event with agency_amount/agent_amount
    populated, not just one flat figure.
    """
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 95001, "Customer ID": "AWCUST1", "Customer Name": "AW Test Customer",
        "Niche/Tablet Price (RM)": 20000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC108-02", "FCC/Agent": "AW Agent",
    }])

    db_path = _db_path(tmp_path)
    result = process_upload(db_path, str(xlsx_path), run_date=today)

    assert len(result["raised_events"]) == 1
    event = result["raised_events"][0]
    # Net Price 20000: agency 7% = 1400.00, agent 8% = 1600.00, total 3000.00 (=15%, same as flat)
    assert event["agency_amount"] == 1400.00
    assert event["agent_amount"] == 1600.00
    assert event["amount"] == 3000.00

    conn = get_connection(db_path)
    row = conn.execute(
        "SELECT agency_amount, agent_amount, amount FROM commission_events WHERE po_no = 95001"
    ).fetchone()
    assert row["agency_amount"] == 1400.00
    assert row["agent_amount"] == 1600.00
    assert row["amount"] == 3000.00


def test_fb_lead_flag_on_contract_applies_the_deduction_at_calculation_time(tmp_path):
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 95002, "Customer ID": "AWCUST2", "Customer Name": "AW Test Customer 2",
        "Niche/Tablet Price (RM)": 20000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC108-02", "FCC/Agent": "AW Agent",
    }])

    db_path = _db_path(tmp_path)

    # Manually flag as FB-lead referred BEFORE the run processes it -
    # simulating a staff member ticking the box in the web app ahead
    # of the cycle that will raise this commission.
    conn = get_connection(db_path)
    from app.db.connection import init_db
    init_db(db_path)
    conn = get_connection(db_path)
    conn.execute(
        "INSERT INTO contracts (po_no, fb_lead_referred, net_price, status) "
        "VALUES (95002, 1, 20000, 'active') "
        "ON CONFLICT(po_no) DO UPDATE SET fb_lead_referred = 1"
    )
    conn.commit()
    conn.close()

    result = process_upload(db_path, str(xlsx_path), run_date=today)

    event = result["raised_events"][0]
    # Agency 7% - 3% deduction = 4% of 20000 = 800.00; agent unaffected at 1600.00
    assert event["agency_amount"] == 800.00
    assert event["agent_amount"] == 1600.00
    assert event["amount"] == 2400.00


def _find_table_rows(sheet, header_marker="PO No"):
    header_row_num = None
    for row in sheet.iter_rows():
        if any(cell.value == header_marker for cell in row):
            header_row_num = row[0].row
            break
    assert header_row_num is not None, f"No header row found on sheet {sheet.title}"
    headers = [cell.value for cell in sheet[header_row_num]]
    data_rows = []
    for row in sheet.iter_rows(min_row=header_row_num + 1):
        values = [cell.value for cell in row]
        if values[headers.index("No")] is None:
            break
        data_rows.append(dict(zip(headers, values)))
    return headers, data_rows


def test_agent_sheet_never_gets_the_split_table_and_shows_only_their_own_cut(tmp_path):
    """
    Regression test: an individual agent's own sheet (e.g. "AW Agent
    X", not the combined "AW Consultancy" sheet) used to ALSO get the
    Agency/Agent split table, and its base commission columns showed
    the full agency+agent combined total rather than just that agent's
    own cut. Confirmed with the business: the split only ever computes
    on the agency-level view; an agent's own sheet just shows what's
    owed to them.
    """
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 95010, "Customer ID": "AWCUST10", "Customer Name": "AW Test Customer 10",
        "Niche/Tablet Price (RM)": 20000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC108-02", "FCC/Agent": "AW Agent X",
    }])

    db_path = _db_path(tmp_path)
    result = process_upload(db_path, str(xlsx_path), run_date=today)
    confirm_all_pending(db_path, result["commission_run_id"])

    report_path = tmp_path / "report.xlsx"
    generate_report(db_path, result["commission_run_id"], str(report_path))

    workbook = openpyxl.load_workbook(report_path)
    assert "AW Agent X" in workbook.sheetnames
    agent_sheet = workbook["AW Agent X"]

    all_values = [cell.value for row in agent_sheet.iter_rows() for cell in row]
    assert "Full Payment Commissioin (RM)" not in all_values  # no split table on the agent's own sheet

    _, agent_rows = _find_table_rows(agent_sheet)
    # Net Price 20000: agent's own cut at 8% full payment = 1600.00 -
    # NOT the combined agency+agent total of 3000.00 (7% + 8%).
    assert agent_rows[0]["Full Payment Commission (RM)"] == 1600.00

    # The combined AW Consultancy sheet still shows the full total.
    group_sheet = workbook["AW Consultancy"]
    _, group_rows = _find_table_rows(group_sheet)
    assert group_rows[0]["Full Payment Commission (RM)"] == 3000.00


def test_flat_agency_still_has_no_split_amounts(tmp_path):
    """Regression check: AC001 (flat) must not suddenly get split amounts."""
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 95003, "Customer ID": "FLATCUST1", "Customer Name": "Flat Test Customer",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001", "FCC/Agent": "Flat Agent",
    }])

    db_path = _db_path(tmp_path)
    result = process_upload(db_path, str(xlsx_path), run_date=today)

    event = result["raised_events"][0]
    assert event["agency_amount"] is None
    assert event["agent_amount"] is None
    assert event["amount"] == 1500.00  # 15% flat, unaffected by any of the split logic
