"""
Step 4 verification: the web layer (login, upload, results, download).

Run with: pytest tests/test_web.py -v
"""

import datetime
import io
import os
import re

import openpyxl
import pytest

from app.db.connection import get_connection
from app.web import create_app
from app.web.auth import hash_password, verify_password
from tests.helpers import build_aor_report, build_master_report


def test_verify_password_with_no_stored_hash_returns_false_not_a_crash():
    """
    users.password_hash is nullable in the schema - verify_password
    must fail closed on a missing hash, not raise.
    """
    assert verify_password("anything", None) is False
    assert verify_password("anything", "") is False


@pytest.fixture
def app(tmp_path):
    db_path = str(tmp_path / "ledger.db")
    application = create_app(db_path, secret_key="test-secret-not-for-real-use")
    application.config["TESTING"] = True

    conn = get_connection(db_path)
    conn.execute(
        "INSERT INTO users (email, password_hash, display_name) VALUES (?, ?, ?)",
        ("staff@xekl.com", hash_password("correct-horse-battery"), "Staff Member"),
    )
    conn.commit()
    conn.close()

    return application


@pytest.fixture
def client(app):
    return app.test_client()


def _csrf_token(client, get_path):
    """Fetches a page and pulls the csrf_token hidden field out of it -
    every real form submission needs a token from a page the same
    session actually rendered."""
    page = client.get(get_path)
    match = re.search(rb'name="csrf_token" value="([^"]+)"', page.data)
    assert match is not None, f"No csrf_token field found on {get_path}"
    return match.group(1).decode()


def _login(client, email="staff@xekl.com", password="correct-horse-battery"):
    token = _csrf_token(client, "/login")
    return client.post(
        "/login",
        data={"email": email, "password": password, "csrf_token": token},
        follow_redirects=True,
    )


def _upload(client, file_path, filename="upload.xlsx", path="/upload", extra_data=None):
    token = _csrf_token(client, path)
    with open(file_path, "rb") as f:
        data = {"report_file": (f, filename), "csrf_token": token}
        data.update(extra_data or {})
        return client.post(
            path,
            data=data,
            content_type="multipart/form-data",
        )


def _upload_aor(client, file_path, filename="aor.xlsx", period_start="2026-01-01", period_end="2026-12-31"):
    # Defaults to a wide-open period covering every test fixture's
    # dates, matching this route's behavior before period scoping
    # existed - a test that cares about the scoping itself passes its
    # own narrower period_start/period_end.
    return _upload(
        client, file_path, filename=filename, path="/upload-aor",
        extra_data={"period_start": period_start, "period_end": period_end},
    )


def _confirm_all_pending(client, run_id):
    """Loads the review page for a run and confirms every pending
    checkbox on it - the web-layer equivalent of a human ticking every
    box and clicking "Confirm selected"."""
    page = client.get(f"/review/{run_id}")
    token = re.search(rb'name="csrf_token" value="([^"]+)"', page.data).group(1).decode()
    event_ids = [m.decode() for m in re.findall(rb'name="event_id" value="(\d+)"', page.data)]
    return client.post(
        f"/review/{run_id}/confirm",
        data={"csrf_token": token, "event_id": event_ids},
        follow_redirects=True,
    )


def test_upload_page_requires_login(client):
    response = client.get("/upload")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_wrong_password_is_rejected(client):
    token = _csrf_token(client, "/login")
    response = client.post(
        "/login",
        data={"email": "staff@xekl.com", "password": "wrong-password", "csrf_token": token},
    )
    assert response.status_code == 401
    assert b"Incorrect" in response.data


def test_unknown_email_is_rejected(client):
    token = _csrf_token(client, "/login")
    response = client.post(
        "/login",
        data={"email": "nobody@xekl.com", "password": "anything", "csrf_token": token},
    )
    assert response.status_code == 401


def test_login_without_a_valid_csrf_token_is_rejected(client):
    response = client.post(
        "/login",
        data={"email": "staff@xekl.com", "password": "correct-horse-battery", "csrf_token": "made-up-token"},
    )
    assert response.status_code == 400


def test_correct_login_reaches_upload_page(client):
    response = _login(client)
    assert response.status_code == 200
    assert b"Upload Commission Base Report" in response.data


def test_logout_blocks_further_access(client):
    _login(client)
    client.get("/logout")
    response = client.get("/upload")
    assert response.status_code == 302


def test_full_flow_upload_review_confirm_and_download_produces_a_real_workbook(client, tmp_path):
    """
    End-to-end through the actual HTTP layer, not by calling the
    pipeline functions directly: log in, upload a fake file, read the
    rendered results page, follow the review link, confirm the
    detected commission, follow the download link, and open the actual
    bytes that come back as a real Excel workbook. Nothing is due, and
    no download link exists, until that confirm step happens.
    """
    _login(client)

    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)
    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50001, "Customer ID": "CUSTWEB1", "Customer Name": "Web Test Customer",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001",
    }])

    response = _upload(client, xlsx_path)

    assert response.status_code == 200
    assert b"1500.00" in response.data  # 15% of Net Price 10,000
    assert b"Full Payment" in response.data  # trigger_labels passed into the template
    assert b"Review &amp; confirm" in response.data
    assert b"Download Excel report" not in response.data  # nothing confirmed yet

    match = re.search(rb"/review/(\d+)", response.data)
    assert match is not None
    run_id = int(match.group(1))

    # Not confirmed yet - downloading now must not hand back a blank
    # or broken file, it should bounce back to the review page.
    premature_download = client.get(f"/download/{run_id}", follow_redirects=True)
    assert premature_download.status_code == 200
    assert b"confirmed" in premature_download.data.lower()

    confirm_response = _confirm_all_pending(client, run_id)
    assert confirm_response.status_code == 200
    assert b"Confirmed 1 commission" in confirm_response.data
    assert b"Download Excel report" in confirm_response.data

    download_response = client.get(f"/download/{run_id}")
    assert download_response.status_code == 200
    assert download_response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = openpyxl.load_workbook(io.BytesIO(download_response.data))
    assert "All" in workbook.sheetnames
    assert "AC001" in workbook.sheetnames


def _void_event(client, run_id, event_id, reason="test mistake"):
    token = _csrf_token(client, f"/review/{run_id}")
    return client.post(
        f"/review/{run_id}/void",
        data={"csrf_token": token, "event_id": str(event_id), "reason": reason},
        follow_redirects=True,
    )


def test_voiding_a_confirmed_event_removes_it_from_the_report(client, tmp_path):
    """
    A mistake caught after confirming (wrong price, confirmed by
    accident) must be reversible: voiding it with a reason takes it off
    every report immediately, and the reason stays visible on the
    review page for the record - it's never just silently deleted.
    """
    _login(client)
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)
    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50002, "Customer ID": "CUSTWEB2", "Customer Name": "Web Test Customer 2",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001",
    }])
    response = _upload(client, xlsx_path)
    run_id = int(re.search(rb"/review/(\d+)", response.data).group(1))
    _confirm_all_pending(client, run_id)

    review_page = client.get(f"/review/{run_id}")
    event_id = int(re.search(rb'name="event_id" value="(\d+)"', review_page.data).group(1))

    void_response = _void_event(client, run_id, event_id, reason="typed the wrong price")
    assert void_response.status_code == 200
    assert b"voided" in void_response.data.lower()
    assert b"typed the wrong price" in void_response.data
    assert b"Download Excel report" not in void_response.data  # nothing confirmed on this run anymore

    # Downloading now must bounce back like a never-confirmed run, not
    # hand back a report that still includes the voided figure.
    download_response = client.get(f"/download/{run_id}", follow_redirects=True)
    assert b"confirmed" in download_response.data.lower()


def test_voiding_without_a_reason_is_rejected(client, tmp_path):
    _login(client)
    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)
    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50003, "Customer ID": "CUSTWEB3", "Customer Name": "Web Test Customer 3",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001",
    }])
    response = _upload(client, xlsx_path)
    run_id = int(re.search(rb"/review/(\d+)", response.data).group(1))
    _confirm_all_pending(client, run_id)

    review_page = client.get(f"/review/{run_id}")
    event_id = int(re.search(rb'name="event_id" value="(\d+)"', review_page.data).group(1))

    response = _void_event(client, run_id, event_id, reason="")
    assert response.status_code == 200
    assert b"reason is required" in response.data.lower()
    assert b"Download Excel report" in response.data  # still confirmed - void was rejected


def test_upload_with_nothing_due_shows_message_and_no_review_link(client, tmp_path):
    _login(client)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50002, "Customer ID": "CUSTWEB2", "Customer Name": "Web Test Customer 2",
        # nothing paid yet
    }])

    response = _upload(client, xlsx_path)

    assert response.status_code == 200
    assert b"Nothing newly due" in response.data
    assert b"Review &amp; confirm" not in response.data


def test_past_reports_page_stays_reachable_after_an_upload_with_nothing_new(client, tmp_path):
    """
    Regression test: once an upload finds nothing newly due, its
    results page has no download link at all (nothing new to review
    that cycle) - before /reports existed, a real, already-confirmed
    report from an earlier upload became completely unreachable the
    moment that happened, with no way back to it anywhere in the app.
    """
    _login(client)

    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)
    xlsx1 = tmp_path / "upload1.xlsx"
    build_master_report(xlsx1, [{
        "No": 1, "PO No": 50010, "Customer ID": "CUSTWEB10", "Customer Name": "Web Test Customer 10",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001",
    }])
    response = _upload(client, xlsx1)
    match = re.search(rb"/review/(\d+)", response.data)
    run_id = int(match.group(1))
    _confirm_all_pending(client, run_id)

    # /reports must not be empty even before any upload finds nothing new -
    reports_page = client.get("/reports")
    assert reports_page.status_code == 200
    assert f"/download/{run_id}".encode() in reports_page.data

    # A second upload with nothing newly due leaves no link of its own,
    # but the confirmed run from before must still be reachable via
    # /reports (and the nav link on every page).
    xlsx2 = tmp_path / "upload2.xlsx"
    build_master_report(xlsx2, [{
        "No": 1, "PO No": 50011, "Customer ID": "CUSTWEB11", "Customer Name": "Web Test Customer 11",
        # nothing paid yet
    }])
    second_response = _upload(client, xlsx2)
    assert b"Past Reports" in second_response.data  # pointed at it directly, not left with a dead end

    reports_page_2 = client.get("/reports")
    assert f"/download/{run_id}".encode() in reports_page_2.data

    download_response = client.get(f"/download/{run_id}")
    assert download_response.status_code == 200


def test_past_reports_surfaces_a_run_left_unconfirmed(client, tmp_path):
    """
    Regression test: a run that raised something newly due but was
    never confirmed used to become completely unreachable the moment
    you navigated away from its own results page - it doesn't qualify
    for the "confirmed reports" table (nothing on it is confirmed yet),
    so there was no link back to /review/<run_id> anywhere. The
    "Awaiting review" section on /reports exists specifically so this
    can't happen.
    """
    _login(client)

    today = datetime.date.today()
    settlement_date = today - datetime.timedelta(days=6)
    xlsx = tmp_path / "upload.xlsx"
    build_master_report(xlsx, [{
        "No": 1, "PO No": 50020, "Customer ID": "CUSTWEB20", "Customer Name": "Web Test Customer 20",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": settlement_date,
        "Agency Code": "AC001",
    }])
    response = _upload(client, xlsx)
    match = re.search(rb"/review/(\d+)", response.data)
    run_id = int(match.group(1))
    # Deliberately never confirmed - simulating navigating away.

    reports_page = client.get("/reports")
    assert reports_page.status_code == 200
    assert b"Awaiting review" in reports_page.data
    assert f"/review/{run_id}".encode() in reports_page.data


def test_confirming_nothing_selected_leaves_it_pending(client, tmp_path):
    _login(client)

    today = datetime.date.today()
    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50004, "Customer ID": "CUSTWEB4", "Customer Name": "Web Test Customer 4",
        "Niche/Tablet Price (RM)": 10000,
        "Full Settlement Paid Date": today - datetime.timedelta(days=6),
        "Agency Code": "AC001",
    }])
    response = _upload(client, xlsx_path)
    run_id = int(re.search(rb"/review/(\d+)", response.data).group(1))

    token = _csrf_token(client, f"/review/{run_id}")
    confirm_response = client.post(
        f"/review/{run_id}/confirm",
        data={"csrf_token": token},  # no event_id selected
        follow_redirects=True,
    )
    assert b"Nothing was confirmed" in confirm_response.data

    download_response = client.get(f"/download/{run_id}", follow_redirects=True)
    assert b"confirmed" in download_response.data.lower()


def test_upload_without_choosing_a_file_shows_a_clear_message(client):
    _login(client)
    token = _csrf_token(client, "/upload")
    response = client.post(
        "/upload", data={"csrf_token": token}, content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"Choose a Commission Base Report file" in response.data


def test_upload_without_a_valid_csrf_token_is_rejected(client, tmp_path):
    _login(client)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 50003, "Customer ID": "CUSTWEB3", "Customer Name": "Web Test Customer 3",
    }])

    with open(xlsx_path, "rb") as f:
        response = client.post(
            "/upload",
            data={"report_file": (f, "upload.xlsx"), "csrf_token": "made-up-token"},
            content_type="multipart/form-data",
        )
    assert response.status_code == 400


def test_malicious_filename_cannot_escape_the_temp_directory(client, tmp_path):
    """
    A filename crafted to look like a path-traversal or absolute-path
    attempt must be sanitized before it's ever used to build a file
    path on disk.
    """
    _login(client)

    xlsx_path = tmp_path / "upload.xlsx"
    build_master_report(xlsx_path, [{
        "No": 1, "PO No": 55001, "Customer ID": "CUSTSEC1", "Customer Name": "Security Test",
    }])

    response = _upload(client, xlsx_path, filename="../../../../tmp/evil_traversal_test.xlsx")

    # The point of this test is simply that the request completes
    # normally (the file gets processed inside the sandboxed temp
    # directory) rather than writing anywhere else on disk.
    assert response.status_code == 200
    assert not os.path.exists("/tmp/evil_traversal_test.xlsx")


def test_corrupted_upload_shows_a_friendly_message_not_a_server_error(client, tmp_path):
    not_really_excel = tmp_path / "fake.xlsx"
    not_really_excel.write_bytes(b"this is not a real xlsx file")

    _login(client)
    response = _upload(client, not_really_excel, filename="fake.xlsx")

    assert response.status_code == 400
    assert b"process this file" in response.data  # avoids the apostrophe, which Jinja2 HTML-escapes to &#39;


def test_downloading_a_nonexistent_run_id_returns_404(client):
    _login(client)
    response = client.get("/download/999999")
    assert response.status_code == 404


def test_upload_size_is_bounded(app):
    assert app.config["MAX_CONTENT_LENGTH"] == 20 * 1024 * 1024


def test_aor_upload_page_requires_login(client):
    response = client.get("/upload-aor")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_aor_upload_without_choosing_a_file_shows_a_clear_message(client):
    _login(client)
    token = _csrf_token(client, "/upload-aor")
    response = client.post(
        "/upload-aor",
        data={"csrf_token": token, "period_start": "2026-08-01", "period_end": "2026-08-31"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"Choose an AOR" in response.data


def test_aor_upload_without_a_period_shows_a_clear_message(client):
    _login(client)
    token = _csrf_token(client, "/upload-aor")
    response = client.post(
        "/upload-aor", data={"csrf_token": token}, content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"date range" in response.data


def test_aor_upload_with_start_after_end_shows_a_clear_message(client):
    _login(client)
    token = _csrf_token(client, "/upload-aor")
    response = client.post(
        "/upload-aor",
        data={"csrf_token": token, "period_start": "2026-08-31", "period_end": "2026-08-01"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"start date is after the end date" in response.data


def test_aor_full_flow_fills_paid_date_and_reaches_review(client, tmp_path):
    """
    End-to-end through the actual HTTP layer: upload a Master report
    with a PO that has no First Instalment Paid Date yet, then upload
    an AOR export with a matching installment-1 receipt for it - the
    results page should show it newly detected, with a review link,
    exactly like an ordinary Master report upload would.
    """
    _login(client)

    xlsx_master = tmp_path / "master.xlsx"
    build_master_report(xlsx_master, [{
        "No": 1, "PO No": 70001, "Customer ID": "CUSTWEBAOR1", "Customer Name": "Web AOR Customer 1",
        "Niche/Tablet Price (RM)": 10000,  # 7.5% = 750.00
        "Agency Code": "AC001",
    }])
    _upload(client, xlsx_master)

    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEB-0001",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "PO No": 70001, "Customer ID": "CUSTWEBAOR1", "Customer Name": "Web AOR Customer 1",
        "Reference No": "TRF 10/08/2026 (INST 01/24)",
    }])
    response = _upload_aor(client, xlsx_aor)

    assert response.status_code == 200
    assert b"1 paid-date" in response.data
    assert b"750.00" in response.data
    assert b"Review &amp; confirm" in response.data

    match = re.search(rb"/review/(\d+)", response.data)
    assert match is not None
    run_id = int(match.group(1))
    confirm_response = _confirm_all_pending(client, run_id)
    assert b"Confirmed 1 commission" in confirm_response.data


def test_aor_upload_with_nothing_new_points_to_past_reports(client, tmp_path):
    _login(client)

    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEB-0002",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "PO No": 99999, "Customer ID": "CUSTWEBAOR2", "Customer Name": "Web AOR Customer 2",
        "Reference No": "HLB 000000 STAMP DUTY",
    }])
    response = _upload_aor(client, xlsx_aor)

    assert response.status_code == 200
    assert b"Nothing newly due this cycle" in response.data
    assert b"Past Reports" in response.data


def test_past_reports_lists_aor_uploads_even_with_nothing_newly_due(client, tmp_path):
    """
    An AOR upload that raises nothing new leaves no commission_runs
    row at all, so it can never show up in the "confirmed runs" table
    above - but the file itself was still uploaded and should stay
    reachable from Past Reports, not disappear the moment its results
    page is left.
    """
    _login(client)

    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEB-0004",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "PO No": 99997, "Customer ID": "CUSTWEBAOR4", "Customer Name": "Web AOR Customer 4",
        "Reference No": "HLB 000000 STAMP DUTY",
    }])
    _upload_aor(client, xlsx_aor, filename="my_aor_export.xlsx")

    page = client.get("/reports")
    assert page.status_code == 200
    assert b"my_aor_export.xlsx" in page.data

    match = re.search(rb"/download-aor-annotated/(\d+)", page.data)
    assert match is not None, "Past Reports should link to the annotated download"
    download_response = client.get(
        match.group(0).decode(),
        query_string={"po_period_start": "2026-01-01", "po_period_end": "2026-12-31"},
    )
    assert download_response.status_code == 200


def test_aor_results_page_links_to_the_annotated_download(client, tmp_path):
    _login(client)

    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEB-0003",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "Purchase Statement Date": datetime.date(2026, 8, 5),
        "PO No": 99998, "Customer ID": "CUSTWEBAOR3", "Customer Name": "Web AOR Customer 3",
        "Reference No": "HLB 000000 STAMP DUTY",
    }])
    response = _upload_aor(client, xlsx_aor)
    assert response.status_code == 200

    match = re.search(rb"/download-aor-annotated/(\d+)", response.data)
    assert match is not None, "Results page should link to the annotated download"
    run_id = int(match.group(1))

    download_response = client.get(
        f"/download-aor-annotated/{run_id}",
        query_string={"po_period_start": "2026-08-01", "po_period_end": "2026-08-31"},
    )
    assert download_response.status_code == 200
    assert download_response.headers["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = openpyxl.load_workbook(io.BytesIO(download_response.data))
    # The original sheet comes through untouched...
    original_sheet = workbook.worksheets[0]
    assert original_sheet.cell(row=23, column=1).value == 1
    # ...its Purchase Statement Date puts it in "Payments (PO Date)"...
    payments_sheet = workbook["Payments (PO Date)"]
    assert payments_sheet.cell(row=2, column=1).value == 1
    # ...but a stamp duty row isn't a full-payment or installment 1/6
    # receipt, so the new "Valid Payments (PO Date)" sheet has no data
    # rows for it.
    valid_sheet = workbook["Valid Payments (PO Date)"]
    assert valid_sheet.cell(row=2, column=1).value is None


def test_download_aor_annotated_requires_login(client):
    response = client.get("/download-aor-annotated/1")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_download_aor_annotated_404s_for_an_unknown_run(client):
    _login(client)
    response = client.get(
        "/download-aor-annotated/999999",
        query_string={"po_period_start": "2026-01-01", "po_period_end": "2026-12-31"},
    )
    assert response.status_code == 404


def _void_aor_trigger(client, po_no, trigger_type="installment_1", reason="wrong PO matched"):
    token = _csrf_token(client, "/reports")
    return client.post(
        "/void-aor-trigger",
        data={"csrf_token": token, "po_no": str(po_no), "trigger_type": trigger_type, "reason": reason},
        follow_redirects=True,
    )


def test_void_aor_trigger_route_requires_login(client):
    response = client.post("/void-aor-trigger", data={"po_no": "1", "trigger_type": "installment_1"})
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_void_aor_trigger_without_a_valid_csrf_token_is_rejected(client):
    _login(client)
    response = client.post(
        "/void-aor-trigger",
        data={
            "csrf_token": "made-up-token", "po_no": "1", "trigger_type": "installment_1",
            "reason": "wrong PO matched",
        },
    )
    assert response.status_code == 400


def test_void_aor_trigger_with_a_non_numeric_po_no_is_rejected(client):
    _login(client)
    response = _void_aor_trigger(client, "not-a-number")
    assert response.status_code == 200
    assert b"PO No must be a number" in response.data


def test_void_aor_trigger_with_an_invalid_trigger_type_is_rejected(client):
    _login(client)
    token = _csrf_token(client, "/reports")
    response = client.post(
        "/void-aor-trigger",
        data={
            "csrf_token": token, "po_no": "80100", "trigger_type": "not_a_real_trigger",
            "reason": "wrong PO matched",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"invalid trigger" in response.data


def test_void_aor_trigger_without_a_reason_is_rejected(client, tmp_path):
    _login(client)
    xlsx_master = tmp_path / "master.xlsx"
    build_master_report(xlsx_master, [{
        "No": 1, "PO No": 70010, "Customer ID": "CUSTWEBVOID1", "Customer Name": "Web Void Customer 1",
        "Niche/Tablet Price (RM)": 10000, "Agency Code": "AC001",
    }])
    _upload(client, xlsx_master)
    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEBVOID-0001",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "PO No": 70010, "Customer ID": "CUSTWEBVOID1", "Customer Name": "Web Void Customer 1",
        "Reference No": "TRF 10/08/2026 (INST 01/24)",
    }])
    _upload_aor(client, xlsx_aor)

    response = _void_aor_trigger(client, 70010, reason="")
    assert response.status_code == 200
    assert b"reason is required" in response.data.lower()


def test_void_aor_trigger_with_nothing_to_void_shows_a_clear_message(client):
    _login(client)
    response = _void_aor_trigger(client, 999999999)
    assert response.status_code == 200
    assert b"no AOR receipt found" in response.data


def test_void_aor_trigger_success_lets_a_corrected_reupload_reapply(client, tmp_path):
    """
    Full round trip through the actual HTTP layer, matching the design
    the feature exists for: a receipt landed on the wrong PO, staff
    void its trigger from Past Reports, then re-upload the (corrected)
    AOR export - the same Acknowledgment Receipt No must be picked up
    again instead of being skipped as already-seen.
    """
    _login(client)
    xlsx_master = tmp_path / "master.xlsx"
    build_master_report(xlsx_master, [{
        "No": 1, "PO No": 70011, "Customer ID": "CUSTWEBVOID2", "Customer Name": "Web Void Customer 2",
        "Niche/Tablet Price (RM)": 10000,  # 7.5% = 750.00
        "Agency Code": "AC001",
    }])
    _upload(client, xlsx_master)

    xlsx_aor = tmp_path / "aor.xlsx"
    build_aor_report(xlsx_aor, [{
        "No": 1, "Acknowledgment Receipt No": "RC-WEBVOID-0002",
        "Acknowledgment Receipt Date": datetime.date(2026, 8, 10),
        "PO No": 70011, "Customer ID": "CUSTWEBVOID2", "Customer Name": "Web Void Customer 2",
        "Reference No": "TRF 10/08/2026 (INST 01/24)",
    }])
    upload_response = _upload_aor(client, xlsx_aor)
    assert b"1 paid-date" in upload_response.data

    void_response = _void_aor_trigger(client, 70011, reason="wrong PO matched")
    assert void_response.status_code == 200
    assert b"voided" in void_response.data.lower()
    assert b"re-upload a corrected AOR file" in void_response.data

    # Re-uploading the same file (as if corrected in Kenjin and
    # re-exported with the same Acknowledgment Receipt No) must apply
    # it again, not silently skip it as already-seen.
    reupload_response = _upload_aor(client, xlsx_aor)
    assert b"1 paid-date" in reupload_response.data
    assert b"750.00" in reupload_response.data
    assert b"Review &amp; confirm" in reupload_response.data


def _create_agency(client, agency_code, name="Test Agency", format_key="xemp"):
    token = _csrf_token(client, "/agencies")
    return client.post(
        "/agencies/new",
        data={"csrf_token": token, "agency_code": agency_code, "name": name, "format_key": format_key},
        follow_redirects=True,
    )


def _edit_agency(client, current_code, new_code=None, name="Edited Agency", format_key="xemp"):
    token = _csrf_token(client, f"/agencies/{current_code}/edit")
    return client.post(
        f"/agencies/{current_code}/edit",
        data={
            "csrf_token": token, "agency_code": new_code if new_code is not None else current_code,
            "name": name, "format_key": format_key,
        },
        follow_redirects=True,
    )


def test_agencies_page_requires_login(client):
    response = client.get("/agencies")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_creating_an_agency_shows_it_on_the_list(client):
    _login(client)
    response = _create_agency(client, "AC400", name="Brand New Agency", format_key="lachesis")
    assert response.status_code == 200
    assert b"added" in response.data.lower()
    assert b"AC400" in response.data
    assert b"Brand New Agency" in response.data
    assert b"Full Payout, Agent Breakdown (Lachesis-style)" in response.data


def test_creating_an_agency_without_a_code_is_rejected(client):
    _login(client)
    response = _create_agency(client, "", name="No Code Agency")
    assert response.status_code == 200
    assert b"Agency Code is required" in response.data


def test_creating_an_agency_without_a_name_is_rejected(client):
    _login(client)
    response = _create_agency(client, "AC401", name="")
    assert response.status_code == 200
    assert b"Name is required" in response.data


def test_creating_an_agency_with_an_invalid_format_is_rejected(client):
    _login(client)
    token = _csrf_token(client, "/agencies")
    response = client.post(
        "/agencies/new",
        data={"csrf_token": token, "agency_code": "AC402", "name": "Bad Format", "format_key": "nonsense"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Choose a valid format" in response.data


def test_creating_an_agency_with_a_duplicate_code_is_rejected(client):
    _login(client)
    _create_agency(client, "AC403", name="First")
    response = _create_agency(client, "AC403", name="Second")
    assert response.status_code == 200
    assert b"already exists" in response.data


def test_agencies_route_without_a_valid_csrf_token_is_rejected(client):
    _login(client)
    response = client.post(
        "/agencies/new",
        data={"csrf_token": "made-up-token", "agency_code": "AC404", "name": "No CSRF", "format_key": "xemp"},
    )
    assert response.status_code == 400


def test_edit_agency_page_requires_login(client):
    response = client.get("/agencies/AC001/edit")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_edit_agency_page_404s_for_an_unknown_code(client):
    _login(client)
    response = client.get("/agencies/DOES-NOT-EXIST/edit")
    assert response.status_code == 404


def test_edit_agency_page_preselects_the_current_format(client):
    _login(client)
    _create_agency(client, "AC405", name="AW-Style Agency", format_key="aw")
    page = client.get("/agencies/AC405/edit")
    assert page.status_code == 200
    assert b'value="aw" selected' in page.data


def test_editing_an_agencys_name_and_format_without_changing_the_code(client):
    _login(client)
    _create_agency(client, "AC406", name="Original Name", format_key="xemp")

    response = _edit_agency(client, "AC406", name="Updated Name", format_key="aw")
    assert response.status_code == 200
    assert b"updated" in response.data.lower()
    assert b"Updated Name" in response.data
    assert b"Agency/Agent Split (AW Consultancy-style)" in response.data


def test_renaming_an_agencys_code_moves_its_contracts(client, app, tmp_path):
    """The core safety property, exercised through the actual HTTP
    layer: an agency's contracts must still be tied to it under its new
    code, not orphaned or left under the old one."""
    _login(client)
    _create_agency(client, "AC407", name="Renaming Agency", format_key="xemp")

    xlsx_master = tmp_path / "master.xlsx"
    build_master_report(xlsx_master, [{
        "No": 1, "PO No": 70090, "Customer ID": "CUSTREN1", "Customer Name": "Rename Customer",
        "Niche/Tablet Price (RM)": 10000, "Agency Code": "AC407",
    }])
    _upload(client, xlsx_master)

    response = _edit_agency(client, "AC407", new_code="AC407-NEW", name="Renaming Agency", format_key="xemp")
    assert response.status_code == 200
    assert b"AC407-NEW" in response.data

    page = client.get("/agencies/AC407-NEW/edit")
    assert page.status_code == 200
    old_page = client.get("/agencies/AC407/edit")
    assert old_page.status_code == 404

    conn = get_connection(app.config["DB_PATH"])
    moved_agency_code = conn.execute("SELECT agency_code FROM contracts WHERE po_no = 70090").fetchone()["agency_code"]
    conn.close()
    assert moved_agency_code == "AC407-NEW"


def test_renaming_an_agency_to_a_code_already_in_use_is_rejected(client):
    _login(client)
    _create_agency(client, "AC408", name="Agency Eight")
    _create_agency(client, "AC409", name="Agency Nine")

    response = _edit_agency(client, "AC408", new_code="AC409", name="Agency Eight")
    assert response.status_code == 200
    assert b"already exists" in response.data


