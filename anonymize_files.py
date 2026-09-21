"""
Temporary utility: anonymizes customer and agent names across the Base
Report, AOR, and "overall commission" reference files so they can be
shared and tested with without exposing real customer/agent data.

Run this locally, on your own machine, against your own real files -
nothing gets uploaded anywhere. Every file you pass gets a new
anonymized copy written to --out-dir; your originals are never
modified.

The SAME real customer/agent gets the SAME fake replacement everywhere
it appears, across every file passed in this one run - so PO No,
dates, prices, agency codes, and Lot No all stay real and
cross-referenceable between files; only the people's names change.

Two passes, on purpose: a real "overall commission" export spells an
individual agent's name directly into a sheet's own TAB NAME and into
a merged title cell ("TAN POH HUI COMMISSION PAYOUT AS AT..."), not
just into the FCC/Agent column - so the mapping has to be fully built
from every header-matched column across every file FIRST, then a
second pass scrubs any occurrence of a known real name out of every
sheet title and every other text cell too (titles, Remarks, anywhere),
not just the columns that named the person directly. Agency/company
names (AW Consultancy, PAJEJU Enterprise, ...) are deliberately left
alone - only individual people get anonymized.

Usage:
    python3 anonymize_files.py base_report.xlsx aor_report.xlsx overall_commission.xlsx --out-dir anonymized

Not part of the web app - safe to delete once you're done testing with
real data and don't need it anymore.
"""

import argparse
import os
import re

import openpyxl
from openpyxl.cell.cell import MergedCell

# Header text (case-insensitive, exact match) that marks a column as
# holding a customer/agent name or ID worth anonymizing. If the
# "overall commission" reference file uses different wording for the
# same thing, add it to the matching set below.
_CUSTOMER_ID_HEADERS = {"customer id"}
_CUSTOMER_NAME_HEADERS = {"customer name"}
_PAYOR_NAME_HEADERS = {"payor name"}
_AGENT_NAME_HEADERS = {"fcc/agent", "agent", "agent name"}

# Which of the fields above are actual people's names worth hunting
# for in free text (sheet titles, Remarks, ...) beyond their own
# column - customer_id is a code, not a name, and scrubbing arbitrary
# codes out of free text risks mangling an unrelated Lot No/PO No/
# Agency Code that happens to share a substring, so it's deliberately
# left out of this second pass.
_FREE_TEXT_SCRUB_FIELDS = ("customer_name", "payor_name", "agent_name")


class _FakeValuePool:
    """Hands out sequential fake values, remembering every original ->
    fake mapping so the same real value always gets the same fake one
    everywhere it appears across every file passed in this run."""

    def __init__(self, template):
        self._template = template  # e.g. "Test Customer {n}" or "CUSTTEST{n:04d}"
        self.mapping = {}

    def get(self, original):
        if original is None or str(original).strip() == "":
            return original
        key = str(original).strip()
        if key not in self.mapping:
            self.mapping[key] = self._template.format(n=len(self.mapping) + 1)
        return self.mapping[key]


def _matches(header, target_set):
    return header is not None and str(header).strip().lower() in target_set


def _find_header_hit_blocks(sheet):
    """Every header row in this sheet's first 30 rows containing any
    target header, as (header_row_num, [(column, field), ...]) - the
    same "search, don't assume a fixed row number" approach the rest
    of the tool uses, since title-block height varies between file
    types."""
    blocks = []
    for row in sheet.iter_rows(min_row=1, max_row=30):
        header_hits = []
        for cell in row:
            if _matches(cell.value, _CUSTOMER_ID_HEADERS):
                header_hits.append((cell.column, "customer_id"))
            elif _matches(cell.value, _CUSTOMER_NAME_HEADERS):
                header_hits.append((cell.column, "customer_name"))
            elif _matches(cell.value, _PAYOR_NAME_HEADERS):
                header_hits.append((cell.column, "payor_name"))
            elif _matches(cell.value, _AGENT_NAME_HEADERS):
                header_hits.append((cell.column, "agent_name"))
        if header_hits:
            blocks.append((row[0].row, header_hits))
    return blocks


def collect_names(path, pools):
    """Read-only first pass: registers every real value found under a
    matched column into `pools`, without writing anything - so the
    mapping is complete (built from ALL files) before any file gets
    scrubbed for stray mentions in free text."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        for header_row_num, header_hits in _find_header_hit_blocks(sheet):
            for data_row in sheet.iter_rows(min_row=header_row_num + 1):
                for col_idx, field in header_hits:
                    cell = sheet.cell(row=data_row[0].row, column=col_idx)
                    if isinstance(cell, MergedCell):
                        continue
                    pools[field].get(cell.value)


def _build_name_pattern(pools):
    """A single case-insensitive regex matching any real name already
    known across _FREE_TEXT_SCRUB_FIELDS, longest names first so e.g.
    "Tan Poh Hui" matches whole rather than partially matching on
    "Tan" from a different person's name. Returns (pattern, lookup) or
    (None, None) if nothing was collected."""
    entries = []
    for field in _FREE_TEXT_SCRUB_FIELDS:
        entries.extend(pools[field].mapping.items())
    if not entries:
        return None, None
    entries.sort(key=lambda kv: -len(kv[0]))
    pattern = re.compile("|".join(re.escape(original) for original, _ in entries), re.IGNORECASE)
    lookup = {original.lower(): fake for original, fake in entries}
    return pattern, lookup


def _scrub_text(text, pattern, lookup):
    if pattern is None or not isinstance(text, str):
        return text
    return pattern.sub(lambda m: lookup[m.group(0).lower()], text)


def anonymize_workbook(path, pools, pattern, lookup):
    """
    Second pass, on the real (not data_only) workbook so it's the copy
    actually saved: replaces every header-matched column's values via
    `pools` (already fully populated by collect_names, so this only
    ever looks up existing mappings, never creates new ones), then
    scrubs every sheet's own tab name and every other text cell for
    any stray occurrence of a known real name.
    """
    workbook = openpyxl.load_workbook(path)
    found_columns = []

    for sheet in workbook.worksheets:
        header_blocks = _find_header_hit_blocks(sheet)
        for header_row_num, header_hits in header_blocks:
            found_columns.append((sheet.title, header_row_num, header_hits))
            for data_row in sheet.iter_rows(min_row=header_row_num + 1):
                for col_idx, field in header_hits:
                    cell = sheet.cell(row=data_row[0].row, column=col_idx)
                    # A cell that's part of a merged range (other than
                    # its top-left anchor) has no independent value to
                    # set - real business spreadsheets merge cells for
                    # formatting all the time, so this is expected, not
                    # an error condition worth stopping the run over.
                    if isinstance(cell, MergedCell):
                        continue
                    cell.value = pools[field].get(cell.value)

        sheet.title = _scrub_text(sheet.title, pattern, lookup)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if isinstance(cell.value, str):
                    cell.value = _scrub_text(cell.value, pattern, lookup)

    return workbook, found_columns


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("files", nargs="+", help="xlsx files to anonymize (Base Report, AOR, overall commission)")
    parser.add_argument("--out-dir", default="anonymized")
    args = parser.parse_args()

    os.makedirs(args.out_dir, exist_ok=True)

    pools = {
        "customer_id": _FakeValuePool("CUSTTEST{n:04d}"),
        "customer_name": _FakeValuePool("Test Customer {n}"),
        "payor_name": _FakeValuePool("Test Payor {n}"),
        "agent_name": _FakeValuePool("Test Agent {n}"),
    }

    for path in args.files:
        collect_names(path, pools)

    pattern, lookup = _build_name_pattern(pools)

    for path in args.files:
        workbook, found_columns = anonymize_workbook(path, pools, pattern, lookup)
        out_path = os.path.join(args.out_dir, os.path.basename(path))
        workbook.save(out_path)

        print(f"\n{path} -> {out_path}")
        if not found_columns:
            print(
                "  WARNING: no recognizable Customer ID / Customer Name / Payor "
                "Name / FCC-Agent column found - nothing anonymized in this "
                "file. Check the header names against the sets at the top of "
                "this script and add yours if they differ."
            )
        for sheet_title, header_row_num, hits in found_columns:
            fields = ", ".join(field for _, field in hits)
            print(f"  sheet '{sheet_title}' row {header_row_num}: anonymized [{fields}]")

    print(
        f"\n{len(pools['customer_id'].mapping)} unique customer ID(s), "
        f"{len(pools['customer_name'].mapping)} unique customer name(s), "
        f"{len(pools['payor_name'].mapping)} unique payor name(s), "
        f"{len(pools['agent_name'].mapping)} unique agent name(s) anonymized "
        f"consistently across all {len(args.files)} file(s) - including any "
        f"stray mention of those names in sheet tab names, title blocks, or "
        f"Remarks, not just their own column."
    )


if __name__ == "__main__":
    main()
