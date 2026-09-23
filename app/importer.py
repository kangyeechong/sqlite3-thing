"""
Reads a Kenjin "Commission Base Report" (Master report) Excel file and
loads it into the ledger database.

The separate AOR (Acknowledgment of Receipt) export is a different
file with its own import path - see app/aor.py - not handled here;
this module actively rejects one if it's uploaded through this path by
mistake (see _AOR_SIGNATURE_HEADER below).
"""

import datetime
import os
import re
from dataclasses import dataclass, field

import openpyxl

from . import commission, parsing, rules

# The real Master sheet has a title block in rows 1-5, headers on row 6,
# and data from row 7 onward - confirmed against the real sample file.
# The importer doesn't hardcode "row 6" though: it searches for the row
# containing "PO No", so a sheet with an extra or missing blank row
# doesn't silently misread everything as garbage.
#
# "Niche/Tablet Price (RM)" is included specifically so this can't also
# match the AOR (Acknowledgment of Receipt) export - a real sample of
# that file shares "No", "PO No", and "Customer ID" too (it's the
# uploaded-payment-receipts list, not a sale-price sheet), and used to
# pass this exact check, getting silently imported as if every AOR
# receipt were its own PO - see _AOR_SIGNATURE_HEADER below for the
# other half of that fix (a specific, actionable error instead of a
# corrupted ledger).
_REQUIRED_HEADERS = ("No", "PO No", "Customer ID", "Niche/Tablet Price (RM)")

# A header this specific to the AOR export - confirmed from a real
# sample, never present in a genuine Commission Base Report - used only
# to give a clear, specific error when someone uploads the wrong file
# here (see import_master_report), rather than a generic "couldn't find
# a header row" message that doesn't explain what actually went wrong.
_AOR_SIGNATURE_HEADER = "Acknowledgment Receipt No"

# The trailing "Date Record" summary table's data rows look like
# "As at 05/06/2026" - confirmed against the real file. Matched with a
# regex (not a fixed column position) since where this table starts
# horizontally shifts with the layout.
_DATE_RECORD_PATTERN = re.compile(r"^As at (\d{2})/(\d{2})/(\d{4})$")


@dataclass
class ReviewFlag:
    po_no: object
    check: str
    message: str


@dataclass
class ImportResult:
    # Scoped to PO Nos genuinely new to the ledger this upload, not
    # every row literally in the file - the real Kenjin export is
    # cumulative (an "August" export re-lists every PO back to
    # whenever records began, not just August's new ones), and staff
    # want each upload's results to read as just that cycle's own
    # activity, not the same old months over and over. Every row in
    # the file is still upserted into the ledger regardless (safe and
    # idempotent, and it lets a genuine correction - e.g. a previously
    # cancelled PO turning out to be real after all - still go through
    # normally) - only what's counted and flagged here is scoped.
    contracts_seen: int = 0
    contracts_new: int = 0
    contracts_updated: int = 0
    review_flags: list = field(default_factory=list)
    # How many rows of the sheet's own trailing "Date Record" history
    # were newly imported this time - 0 on every upload after the
    # first (each date only ever gets imported once, see
    # historical_summary_rows.date_record's UNIQUE constraint), not an
    # error or something to re-derive.
    historical_rows_imported: int = 0
    # How many PO numbers were missing from the sequence between the
    # lowest and highest PO No seen in this upload - each one gets a
    # synthetic cancelled contract row (see _detect_cancelled_po_gaps)
    # so it's tracked instead of silently vanishing. Not counted in
    # contracts_seen/contracts_new, since those describe what was new
    # in the uploaded file, not something inferred from it.
    cancelled_po_gaps_detected: int = 0
    # The PO Date range spanned by this upload's genuinely-new rows
    # (ISO date strings), or None if none of them have a PO Date on
    # file at all. Lets the results page offer a one-click link straight
    # to this month's Overall Commission report (see
    # app.report.generate_period_report) without staff having to type
    # dates themselves - the AOR upload still needs a manual period
    # picker (its export is fragmented across arbitrary windows that
    # don't line up with calendar months), but a Master report upload's
    # own new rows always share one real, known purchase-date range.
    new_po_date_min: object = None
    new_po_date_max: object = None


def _find_header_row(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=30):
        values = [cell.value for cell in row]
        if all(h in values for h in _REQUIRED_HEADERS):
            return row[0].row
    raise ValueError(
        f"Could not find a header row containing {_REQUIRED_HEADERS} "
        f"in sheet '{sheet.title}' - is this really a Commission Base "
        f"Report?"
    )


def _is_master_shaped(sheet):
    try:
        _find_header_row(sheet)
        return True
    except ValueError:
        return False


def _looks_like_an_aor_export(sheet):
    """True if this sheet carries the AOR export's own distinctive
    header, whether or not it's Master-shaped - used only by
    import_master_report to give a specific, actionable error when the
    wrong file gets uploaded here, instead of a generic "couldn't find
    a header row" message that doesn't say what actually went wrong."""
    for row in sheet.iter_rows(min_row=1, max_row=30):
        if _AOR_SIGNATURE_HEADER in [cell.value for cell in row]:
            return True
    return False


def _find_date_record_header(sheet):
    """
    Locates the "DATE RECORD" cell of the sheet's own trailing summary
    table - confirmed against the real file, it sits well to the right
    of the PO columns (around column K), not at column A, and its
    exact column shifts with the layout, so every cell is checked
    rather than assuming a fixed position. Returns (row, column), or
    (None, None) if this sheet has no such table (a smaller test
    fixture, for instance - not every upload will have one).
    """
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == "DATE RECORD":
                return cell.row, cell.column
    return None, None


def _read_historical_summary_rows(sheet):
    """
    Reads the sheet's own trailing "Date Record" summary table - the
    permanent record of every processing cycle that happened before
    this tool existed. Returns a list of dicts (date_record as an ISO
    date string, the three commission figures, remarks) in whatever
    order they appear on the sheet.

    Deliberately does NOT read the Running Total column - that number
    is always recomputed fresh from whatever's actually on file (see
    report._load_summary_rows), never trusted as a cached figure that
    could go stale.
    """
    header_row, header_col = _find_date_record_header(sheet)
    if header_row is None:
        return []

    rows = []
    row_num = header_row + 1
    while True:
        label = sheet.cell(row=row_num, column=header_col).value
        if not isinstance(label, str):
            break
        match = _DATE_RECORD_PATTERN.match(label.strip())
        if not match:
            break  # e.g. the "Total Sum of Commission Payout..." row right after
        day, month, year = match.groups()
        rows.append({
            "date_record": datetime.date(int(year), int(month), int(day)).isoformat(),
            "full_commission": _to_number(sheet.cell(row=row_num, column=header_col + 1).value) or 0.0,
            "first_half_commission": _to_number(sheet.cell(row=row_num, column=header_col + 2).value) or 0.0,
            "second_half_commission": _to_number(sheet.cell(row=row_num, column=header_col + 3).value) or 0.0,
            "remarks": sheet.cell(row=row_num, column=header_col + 5).value,
        })
        row_num += 1
    return rows


def _existed_by_cutoff(contract, cutoff_iso_date):
    """
    True if this contract's own PO Date or Signature Date - a real,
    immutable fact about when the sale actually happened, not when OUR
    ledger happened to first import its row - falls on or before the
    cutoff. Prefers Signature Date (when the customer actually signed,
    a firmer signal the sale was final) and falls back to PO Date.

    This is what _flag_historically_accounted_commissions uses to
    decide whether a PO could possibly be money a historical cutoff
    already accounts for, instead of "is this PO new to my local
    contracts table" (a database-instance-local signal, not a business
    fact - see the regression this replaced). A local database reset
    (contracts cleared but historical_summary_rows/agencies left
    alone, or a fresh ledger re-onboarding a company's full existing
    book) makes every real, long-standing PO look "new to the ledger"
    even though it may have existed for months - treating that as
    equivalent to a genuinely new sale would wrongly re-detect the
    company's entire historical commission total as newly due all over
    again the moment the ledger is rebuilt.

    Neither date on file is the safe direction to fail in: it means
    this PO falls through to normal live detection instead of being
    silently absorbed - the same "prefer over-detecting to silently
    losing money" reasoning used throughout this module.
    """
    sale_date = contract["signature_date"] or contract["po_date"]
    return bool(sale_date) and sale_date <= cutoff_iso_date


def _flag_historically_accounted_commissions(conn, cutoff_iso_date):
    """
    Runs on every upload that has a trailing Date Record history at
    all (see import_master_report) - marks every trigger that would
    have legitimately been due on or before that history's own last
    "As at" date as already flagged, WITHOUT creating a
    commission_event for it.

    Why: that money is already counted in the aggregate historical
    total. Without this, process_commission_run (which runs right
    after, in the same upload) would see these same PO's paid-dates
    and detect them as newly due too - double-counting real money. A
    paid-date strictly AFTER the cutoff is left alone and still goes
    through normal detection, so a genuinely new payment isn't
    silently swallowed by it. Re-running this on every upload (not
    just the one that first imports the history) matters because a PO
    that already existed before the cutoff was established can still
    have its paid-date field filled in on a *later* upload - that
    payment is just as much already covered by the historical total as
    one whose date was already on file the first time.

    Every candidate must also pass _existed_by_cutoff: a PO whose own
    sale date is AFTER the cutoff cannot possibly be money that
    cutoff's total already accounts for, no matter what its paid-date
    says - flagging it here instead of letting it go through normal
    detection would silently and permanently lose that commission (it
    would never get a commission_event, and a flag is never unset
    anywhere in this codebase). Regression case: an upload introduces a
    PO that never appeared before, with a paid-date that happens to
    predate an already-established cutoff (e.g. a late Kenjin entry for
    an older sale) - without this check, that PO's flag gets set here
    and the commission never surfaces for review at all.

    Deliberately reuses full_payment_is_due/installment_1_is_due/
    installment_6_is_due (the exact same eligibility checks
    process_commission_run itself uses) rather than a simpler
    date-only check: those functions already refuse to flag a contract
    with a non-positive Net Price precisely so that once staff fix the
    underlying data, it's still eligible to be flagged correctly later
    (see commission.py) - a flag is permanent and never gets unset
    anywhere in this codebase, so setting one from bad or
    not-yet-active data here would silently and permanently lose that
    commission even after the data is corrected. Same reasoning covers
    a cancelled/on_hold contract (status != 'active') and an At-Need
    case still missing its inurnment date.
    """
    cutoff_date = datetime.date.fromisoformat(cutoff_iso_date)
    candidates = conn.execute("SELECT * FROM contracts WHERE status = 'active'").fetchall()
    for contract in candidates:
        if not _existed_by_cutoff(contract, cutoff_iso_date):
            continue
        # full_payment_is_due's At-Need branch deliberately ignores
        # `as_of` entirely (there's no cooling-off wait for At-Need -
        # see commission.py) - it only checks that an inurnment date
        # is on file, regardless of when. That's correct for live
        # detection (as_of is always "today"), but here as_of is a
        # past cutoff date, so an explicit settlement-date check is
        # still needed on top of it for At-Need - otherwise a
        # genuinely new At-Need payment dated AFTER the cutoff would
        # get permanently flagged as "already accounted for" too.
        if (
            contract["full_settlement_paid_date"]
            and contract["full_settlement_paid_date"] <= cutoff_iso_date
            and commission.full_payment_is_due(contract, cutoff_date)
        ):
            conn.execute(
                "UPDATE contracts SET full_commission_flagged = 1 WHERE po_no = ?",
                (contract["po_no"],),
            )
        if (
            contract["first_installment_paid_date"]
            and contract["first_installment_paid_date"] <= cutoff_iso_date
            and commission.installment_1_is_due(contract)
        ):
            conn.execute(
                "UPDATE contracts SET installment_1_commission_flagged = 1 WHERE po_no = ?",
                (contract["po_no"],),
            )
        if (
            contract["sixth_installment_paid_date"]
            and contract["sixth_installment_paid_date"] <= cutoff_iso_date
            and commission.installment_6_is_due(contract)
        ):
            conn.execute(
                "UPDATE contracts SET installment_6_commission_flagged = 1 WHERE po_no = ?",
                (contract["po_no"],),
            )


def _to_iso_date(value):
    """Normalizes an Excel date cell (datetime, date, or None) to an
    ISO date string, or None."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return None


def _to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _is_positive_whole_number(value):
    """
    True for 1, 2, 3... whether Excel/openpyxl hands it back as an int
    or as a float (a "No" or "PO No" column with General number format
    can come back as 1.0 rather than 1) - checked explicitly because
    isinstance(value, int) alone would silently misread a float-typed
    column as the end of the data and truncate the import.
    """
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return value > 0
    if isinstance(value, float):
        return value > 0 and value.is_integer()
    return False


def _read_rows(sheet):
    """
    Yields one dict per data row, keyed by the sheet's own header text.
    Stops as soon as a row's "No" column isn't a positive integer -
    that's how the real sheet's trailing summary table (Date Record /
    Full Commission / ... - see docs/data_model.md section 2) gets
    excluded without needing to know its exact row position.
    """
    header_row_num = _find_header_row(sheet)
    headers = [cell.value for cell in sheet[header_row_num]]

    for row in sheet.iter_rows(min_row=header_row_num + 1):
        values = [cell.value for cell in row]
        row_no = values[headers.index("No")]
        if not _is_positive_whole_number(row_no):
            break
        yield dict(zip(headers, values))


def _build_contract_fields(raw_row):
    niche_price = _to_number(raw_row.get("Niche/Tablet Price (RM)")) or 0.0
    promotion = _to_number(raw_row.get("Promotion (RM)")) or 0.0
    discount = _to_number(raw_row.get("Discount (RM)")) or 0.0
    net_price = round(niche_price - promotion - discount, 2)

    remarks = raw_row.get("Remarks")

    return {
        "po_no": int(raw_row["PO No"]),
        "customer_id": raw_row.get("Customer ID"),
        "customer_name": raw_row.get("Customer Name"),
        "agent_name": raw_row.get("FCC/Agent"),
        "agency_code": raw_row.get("Agency Code") or None,
        "lot_no": raw_row.get("Lot No"),
        "po_date": _to_iso_date(raw_row.get("PO Date")),
        "signature_date": _to_iso_date(raw_row.get("Signature Date")),
        "niche_price": niche_price,
        "promotion": promotion,
        "discount": discount,
        "net_price": net_price,
        "case_type": parsing.detect_case_type(remarks),
        "inurnment_date": parsing.extract_inurnment_date(remarks),
        "status": parsing.detect_status(remarks),
        "full_settlement_paid_date": _to_iso_date(raw_row.get("Full Settlement Paid Date")),
        "first_installment_paid_date": _to_iso_date(raw_row.get("First Instalment Paid Date")),
        "sixth_installment_paid_date": _to_iso_date(raw_row.get("Sixth Instalment Paid Date")),
        # Accounts fills these in by hand once they've actually sent the
        # money - this tool only ever reads them back, never writes them.
        "full_commission_paid_date": _to_iso_date(raw_row.get("Full Commission Paid Date")),
        "installment_1_commission_paid_date": _to_iso_date(raw_row.get("1st Half Commission Paid Date")),
        "installment_6_commission_paid_date": _to_iso_date(raw_row.get("Balance Half Commission Paid Date")),
        "remarks": remarks,
    }


def _review_checks(conn, fields_list, already_known_po_nos):
    """
    Scans the whole batch for anomalies before anything is written.
    Returns a list of ReviewFlag - informational only, nothing here
    blocks the import. See docs/data_model.md section 6b for the
    rationale behind each check.

    already_known_po_nos: PO Nos the ledger already had real data for
    before this upload (see import_master_report). Every check below
    except status_changed is a pure per-row data-quality check with no
    memory of history - re-running it on a PO the ledger already knows
    about, whose data hasn't changed, would just re-flag the exact same
    already-seen issue every time a cumulative export repeats an old
    month, which is pure noise to a human skimming results. Those
    checks are skipped for an already-known PO. status_changed is the
    one exception: it's inherently comparative (only ever fires when
    Remarks-derived status genuinely differs from what's on file), so
    it's checked for every row regardless of scope - it naturally stays
    silent for an unchanged already-known row on its own.
    """
    flags = []
    seen_po_nos = set()
    known_agency_codes = {
        row["agency_code"] for row in conn.execute("SELECT agency_code FROM agencies")
    }

    for fields in fields_list:
        po_no = fields["po_no"]

        if po_no not in already_known_po_nos:
            if po_no in seen_po_nos:
                flags.append(ReviewFlag(po_no, "duplicate_po", "This PO No appears more than once in this upload."))
            seen_po_nos.add(po_no)

            if fields["first_installment_paid_date"] and fields["sixth_installment_paid_date"]:
                if fields["sixth_installment_paid_date"] < fields["first_installment_paid_date"]:
                    flags.append(ReviewFlag(
                        po_no, "dates_out_of_order",
                        "Sixth Instalment Paid Date is earlier than First Instalment Paid Date.",
                    ))

            if fields["net_price"] is not None and fields["net_price"] <= 0:
                flags.append(ReviewFlag(po_no, "non_positive_net_price", f"Net Price is {fields['net_price']}."))

            if fields["agency_code"] and fields["agency_code"] not in known_agency_codes:
                is_aw_code = fields["agency_code"] in rules.AGENCIES_WITH_AGENCY_AGENT_SPLIT
                flags.append(ReviewFlag(
                    po_no, "unknown_agency_code",
                    f"Agency Code '{fields['agency_code']}' hasn't been seen before - check for a typo. "
                    + (
                        "Recognized as an AW Consultancy code, so it'll get the agency/agent split."
                        if is_aw_code else
                        "New agencies default to flat commission (no agency/agent split) unless "
                        "added to rules.AGENCIES_WITH_AGENCY_AGENT_SPLIT - if this is actually another "
                        "AW Consultancy agent code, add it there before relying on this run's numbers."
                    ),
                ))
                known_agency_codes.add(fields["agency_code"])  # don't re-flag within the same batch

            if fields["status"] == "active" and not fields["agency_code"]:
                flags.append(ReviewFlag(po_no, "blank_agency_code", "Agency Code is blank on an active PO."))

            has_any_payment = any([
                fields["full_settlement_paid_date"],
                fields["first_installment_paid_date"],
                fields["sixth_installment_paid_date"],
            ])
            if fields["status"] in ("cancelled", "withdrawn", "on_hold") and has_any_payment:
                flags.append(ReviewFlag(
                    po_no, "payment_on_inactive_po",
                    f"Status is '{fields['status']}' but a paid date is present.",
                ))

            if fields["case_type"] == "at_need" and fields["inurnment_date"] is None:
                flags.append(ReviewFlag(
                    po_no, "at_need_missing_inurnment_date",
                    "Remarks mention an At Need case but no inurnment date could be read from it.",
                ))

        existing = conn.execute(
            "SELECT status FROM contracts WHERE po_no = ?", (po_no,)
        ).fetchone()
        if existing is not None and existing["status"] != fields["status"]:
            flags.append(ReviewFlag(
                po_no, "status_changed",
                f"Status changed from '{existing['status']}' to '{fields['status']}' "
                f"based on Remarks text - please confirm this is correct.",
            ))

    return flags


def _upsert_contract(conn, fields, now_iso):
    # Guarded the same way the agency_code insert below is: customers.
    # customer_id is TEXT PRIMARY KEY, and SQLite's uniqueness check
    # never treats one NULL as equal to another, so
    # "ON CONFLICT(customer_id)" never fires for a NULL customer_id -
    # every blank-customer row (most visibly, every synthetic
    # cancelled-PO placeholder from _detect_cancelled_po_gaps, which
    # never has a customer at all) would otherwise insert its own new
    # junk row into customers instead of being caught by the conflict
    # clause, growing that table without bound across uploads.
    if fields["customer_id"]:
        conn.execute(
            "INSERT INTO customers (customer_id, name) VALUES (?, ?) "
            "ON CONFLICT(customer_id) DO UPDATE SET name = excluded.name",
            (fields["customer_id"], fields["customer_name"]),
        )

    if fields["agency_code"]:
        splits_by_agent = 0 if fields["agency_code"] in rules.AGENCIES_WITHOUT_PER_AGENT_SPLIT else 1
        commission_split_type = (
            "agency_agent_split" if fields["agency_code"] in rules.AGENCIES_WITH_AGENCY_AGENT_SPLIT else "flat"
        )
        agency_group = rules.AGENCY_GROUPS.get(fields["agency_code"])
        conn.execute(
            "INSERT INTO agencies (agency_code, splits_by_agent, commission_split_type, agency_group) "
            "VALUES (?, ?, ?, ?) "
            "ON CONFLICT(agency_code) DO NOTHING",
            (fields["agency_code"], splits_by_agent, commission_split_type, agency_group),
        )

    existing = conn.execute(
        "SELECT po_no FROM contracts WHERE po_no = ?", (fields["po_no"],)
    ).fetchone()

    conn.execute(
        """
        INSERT INTO contracts (
            po_no, customer_id, agent_name, agency_code, lot_no,
            po_date, signature_date, niche_price, promotion, discount,
            net_price, case_type, inurnment_date, status,
            full_settlement_paid_date, first_installment_paid_date,
            sixth_installment_paid_date,
            full_commission_paid_date, installment_1_commission_paid_date,
            installment_6_commission_paid_date,
            remarks, updated_at
        ) VALUES (
            :po_no, :customer_id, :agent_name, :agency_code, :lot_no,
            :po_date, :signature_date, :niche_price, :promotion, :discount,
            :net_price, :case_type, :inurnment_date, :status,
            :full_settlement_paid_date, :first_installment_paid_date,
            :sixth_installment_paid_date,
            :full_commission_paid_date, :installment_1_commission_paid_date,
            :installment_6_commission_paid_date,
            :remarks, :now
        )
        ON CONFLICT(po_no) DO UPDATE SET
            customer_id = excluded.customer_id,
            agent_name = excluded.agent_name,
            agency_code = excluded.agency_code,
            lot_no = excluded.lot_no,
            po_date = excluded.po_date,
            signature_date = excluded.signature_date,
            niche_price = excluded.niche_price,
            promotion = excluded.promotion,
            discount = excluded.discount,
            net_price = excluded.net_price,
            case_type = excluded.case_type,
            inurnment_date = excluded.inurnment_date,
            status = excluded.status,
            -- COALESCE(excluded.x, contracts.x), not a blind overwrite:
            -- the real Master report never actually carries these three
            -- dates (confirmed against real data - every row's blank),
            -- Accounts fills them in by hand or via an AOR upload
            -- instead. A blind overwrite here would silently wipe a
            -- paid-date an AOR upload already confirmed back to NULL
            -- the next time that same PO appears in a re-uploaded
            -- Master report - found via reproduction, not theoretical.
            -- The commission itself is never double-counted either way
            -- (see the *_commission_flagged columns in commission.py),
            -- but the wiped date would still show blank on a report
            -- generated afterwards even though it was truly paid.
            full_settlement_paid_date = COALESCE(excluded.full_settlement_paid_date, contracts.full_settlement_paid_date),
            first_installment_paid_date = COALESCE(excluded.first_installment_paid_date, contracts.first_installment_paid_date),
            sixth_installment_paid_date = COALESCE(excluded.sixth_installment_paid_date, contracts.sixth_installment_paid_date),
            -- Same COALESCE reasoning as the three paid-date columns
            -- above, and for the same underlying reason: these three
            -- are also hand-typed onto the real Master Report by
            -- Accounts once they've actually sent the money (see the
            -- schema.sql comment on these columns), not something
            -- Kenjin's own export retains on a fresh re-generation. A
            -- blind overwrite here has the exact same failure mode the
            -- three paid-date columns had - a later Master report
            -- re-upload for a PO Accounts already marked as commission-
            -- paid would silently wipe that paid-date back to NULL,
            -- making an already-paid commission look unpaid on the next
            -- downloaded report.
            full_commission_paid_date = COALESCE(excluded.full_commission_paid_date, contracts.full_commission_paid_date),
            installment_1_commission_paid_date = COALESCE(excluded.installment_1_commission_paid_date, contracts.installment_1_commission_paid_date),
            installment_6_commission_paid_date = COALESCE(excluded.installment_6_commission_paid_date, contracts.installment_6_commission_paid_date),
            remarks = excluded.remarks,
            updated_at = excluded.updated_at
        """,
        {**fields, "now": now_iso},
    )

    return existing is None


# A gap this large is far more likely to be a data-entry mistake (an
# extra digit typed into one PO No) than several hundred consecutive
# real cancellations - past this size, _detect_cancelled_po_gaps stops
# and flags the whole range for a human instead of trying to insert
# that many synthetic rows.
_MAX_AUTO_FILLED_GAP = 500


def _detect_cancelled_po_gaps(conn, po_nos_this_upload, now_iso):
    """
    A PO No missing from the sequence between the lowest and highest
    number seen in this upload always means that PO was cancelled
    before it was ever finalized - confirmed with the business (never
    a reserved-but-unused number, a different office's own numbering,
    or simply a PO that hasn't reached this file yet - that last case
    is exactly why this is scoped to the min...max range actually seen
    THIS upload, never extrapolated past the highest number: a number
    beyond that hasn't happened yet, that's not a gap).

    Every missing number not already a real contract (from this or any
    earlier upload) gets a synthetic placeholder row - no customer, no
    price, status 'cancelled', Remarks "Cancelled PO" - built the exact
    same way a real row would be (via _build_contract_fields, so
    detect_status/case_type/etc. all run identically) so it shows up
    in the report exactly like any other cancelled PO: a beige row,
    excluded from commission detection, instead of just disappearing.

    If a real row for one of these numbers shows up on a later upload
    (the business un-cancels it, or this was a mistake), the normal
    upsert path overwrites this placeholder with the real data, same
    as updating any other existing PO - nothing special needed for
    that to work correctly.

    A placeholder has no PO Date of its own - Kenjin never assigned it
    one, since the PO was never actually finalized. But PO Nos are
    handed out in roughly the order contracts get signed, so the
    nearest real PO No's own PO Date (within this same upload's range)
    is a solid stand-in, and it's not just cosmetic: the Overall
    Commission period report is scoped by po_date (see
    generate_period_report), so a placeholder left with no po_date at
    all would silently vanish from every period download forever, even
    though it still shows correctly on the unscoped report. Every
    upload also gets a fresh chance to backfill an older placeholder
    that's still missing a date - from before this inference existed,
    or one that simply had no dated neighbor yet at the time.
    """
    if not po_nos_this_upload:
        return []

    lowest, highest = min(po_nos_this_upload), max(po_nos_this_upload)
    if highest - lowest > _MAX_AUTO_FILLED_GAP:
        return [ReviewFlag(
            None, "po_range_too_wide_to_scan",
            f"PO No ranges from {lowest} to {highest} in this upload - too wide a span to "
            f"safely check for cancelled-PO gaps (likely a typo in one PO No rather than "
            f"{highest - lowest} real cancellations). Skipped; check for a data entry error.",
        )]

    neighbor_rows = conn.execute(
        "SELECT po_no, po_date, customer_id FROM contracts WHERE po_no BETWEEN ? AND ?",
        (lowest, highest),
    ).fetchall()
    already_real = {row["po_no"] for row in neighbor_rows}
    dated_neighbors = {row["po_no"]: row["po_date"] for row in neighbor_rows if row["po_date"]}

    def _nearest_date(po_no):
        if not dated_neighbors:
            return None
        nearest_po_no = min(dated_neighbors, key=lambda n: (abs(n - po_no), n))
        return dated_neighbors[nearest_po_no]

    # customer_id IS NULL is the same signal already_known_po_nos and
    # _upsert_contract rely on to mean "our own synthetic placeholder,
    # never a real contract" - safe to backfill without risking a real
    # PO whose Master report row genuinely has no PO Date.
    for row in neighbor_rows:
        if row["customer_id"] is None and not row["po_date"]:
            inferred = _nearest_date(row["po_no"])
            if inferred:
                conn.execute(
                    "UPDATE contracts SET po_date = ?, updated_at = ? WHERE po_no = ?",
                    (inferred, now_iso, row["po_no"]),
                )

    missing = [n for n in range(lowest, highest + 1) if n not in po_nos_this_upload]
    flags = []
    for po_no in missing:
        if po_no in already_real:
            continue  # already a real contract from an earlier upload - not a gap after all
        fields = _build_contract_fields({"PO No": po_no, "Remarks": "Cancelled PO"})
        fields["po_date"] = _nearest_date(po_no)
        _upsert_contract(conn, fields, now_iso)
        flags.append(ReviewFlag(
            po_no, "inferred_cancelled_po",
            f"PO No {po_no} is missing from the sequence (between {lowest} and {highest} "
            f"in this upload) - added as a cancelled PO so it's tracked, not silently skipped.",
        ))
    return flags


def import_master_report(conn, file_path, imported_by_user=None):
    """
    Reads every sheet in the workbook that looks like a Master report
    (has the expected headers), upserts every PO into `contracts`, and
    returns an ImportResult with counts and review flags.

    Deliberately does not commit the transaction - the caller decides
    when to commit, so a review of the flags can happen before the
    import is finalized if desired.
    """
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    now_iso = datetime.datetime.now().isoformat()

    # Only the Master sheet is read, not the per-agency sheets that
    # follow it in the same workbook - those are filtered views of the
    # identical rows (confirmed against the real sample file), not
    # additional data. Reading every sheet would mean seeing each PO
    # multiple times and having to guess whether a repeat is that
    # expected cross-sheet echo or a genuine duplicate-entry mistake;
    # reading only the first matching sheet avoids the ambiguity
    # entirely, so a real duplicate PO *within* that sheet still gets
    # caught by the review panel below instead of being masked by it.
    master_sheet = next(
        (sheet for sheet in workbook.worksheets if _is_master_shaped(sheet)),
        None,
    )
    if master_sheet is None:
        if any(_looks_like_an_aor_export(sheet) for sheet in workbook.worksheets):
            raise ValueError(
                f"'{file_path}' looks like an Acknowledgment of Receipt (AOR) "
                f"export, not a Commission Base Report - this upload only "
                f"accepts the Kenjin Master report."
            )
        raise ValueError(
            f"No sheet in '{file_path}' has the expected headers "
            f"{_REQUIRED_HEADERS} - is this really a Commission Base Report?"
        )

    all_fields = []
    skip_flags = []
    for raw_row in _read_rows(master_sheet):
        po_no_value = raw_row.get("PO No")
        if not _is_positive_whole_number(po_no_value):
            # Can't import a row with no usable primary key - skip just
            # this row rather than letting int() raise and crash the
            # whole upload for every other valid row in the file.
            skip_flags.append(ReviewFlag(
                raw_row.get("No"), "missing_po_no",
                f"Row 'No'={raw_row.get('No')} has no valid PO No and was not imported.",
            ))
            continue
        all_fields.append(_build_contract_fields(raw_row))

    # Snapshotted before any upsert below runs, so this reflects what
    # the ledger knew BEFORE this upload - see ImportResult's own
    # comment for why this scoping exists. customer_id IS NOT NULL
    # deliberately excludes our own synthetic cancelled-PO gap
    # placeholders (see _detect_cancelled_po_gaps - a placeholder
    # always has no customer, the same signal _upsert_contract's own
    # customers-table guard relies on) - a real row finally arriving
    # for one of those numbers is genuine news (the business un-
    # cancelled it, or the gap-fill was wrong), not an already-known PO
    # repeating itself, so it must still be reported as new/updated.
    already_known_po_nos = {
        row["po_no"] for row in conn.execute("SELECT po_no FROM contracts WHERE customer_id IS NOT NULL")
    }
    new_fields = [f for f in all_fields if f["po_no"] not in already_known_po_nos]

    result = ImportResult(contracts_seen=len(new_fields))
    result.review_flags = skip_flags + _review_checks(conn, all_fields, already_known_po_nos)

    new_po_dates = sorted(f["po_date"] for f in new_fields if f["po_date"])
    if new_po_dates:
        result.new_po_date_min = new_po_dates[0]
        result.new_po_date_max = new_po_dates[-1]

    for fields in all_fields:
        # Always upserted, known-before-this-upload or not - see
        # ImportResult's comment: this keeps every PO's data fresh and
        # lets a real correction through, only the counts below (and
        # review_flags above) are scoped to what's new.
        is_new = _upsert_contract(conn, fields, now_iso)
        if fields["po_no"] in already_known_po_nos:
            continue
        if is_new:
            result.contracts_new += 1
        else:
            result.contracts_updated += 1

    # Deliberately scoped to every PO in the file, not just new_fields:
    # an old month's gap that's already filled in is still correctly
    # excluded by _detect_cancelled_po_gaps's own "already a real
    # contract" check below, so re-scanning the whole file's range
    # here doesn't re-flag or re-create anything - it's already
    # naturally idempotent without needing its own new/known split.
    # Runs after every real row above is already in the database, so a
    # gap only ever means "genuinely missing from this upload" - never
    # a false positive against a PO this same upload was about to add.
    gap_flags = _detect_cancelled_po_gaps(conn, {f["po_no"] for f in all_fields}, now_iso)
    result.review_flags.extend(gap_flags)
    # Excludes the single "too wide to scan" warning flag, which
    # represents zero actual gaps filled, not one.
    result.cancelled_po_gaps_detected = sum(1 for f in gap_flags if f.check == "inferred_cancelled_po")

    # The sheet's own trailing "Date Record" history, imported once and
    # kept forever - see historical_summary_rows in schema.sql. Every
    # date is only ever imported the first time it's seen (UNIQUE
    # constraint, ON CONFLICT DO NOTHING), so re-uploading the same or
    # a later file that repeats these same historical rows never
    # duplicates or overwrites them - they're a permanent fact, not
    # something re-derived on every import.
    historical_rows = _read_historical_summary_rows(master_sheet)
    for historical_row in historical_rows:
        cursor = conn.execute(
            """
            INSERT INTO historical_summary_rows (
                date_record, full_commission, first_half_commission,
                second_half_commission, remarks, imported_at, imported_by_user, source_filename
            ) VALUES (:date_record, :full_commission, :first_half_commission,
                :second_half_commission, :remarks, :now, :imported_by_user, :source_filename)
            ON CONFLICT(date_record) DO NOTHING
            """,
            {
                **historical_row, "now": now_iso,
                "imported_by_user": imported_by_user,
                "source_filename": os.path.basename(file_path),
            },
        )
        if cursor.rowcount:
            result.historical_rows_imported += 1

    # Whatever the sheet's history already accounts for (up to the
    # LATEST "As at" date ever recorded across every upload ever done,
    # not just what's in this specific sheet) must not also get
    # detected as newly due below - that would double-count real
    # money. Read back from the database rather than only this
    # upload's own historical_rows list: an older or stale file (one
    # whose own trailing table hasn't caught up to a later "As at" row
    # a previous upload already established) must not use a smaller
    # cutoff than what's already known, or POs paid in that gap would
    # wrongly fall through to fresh detection.
    if historical_rows:
        cutoff = conn.execute(
            "SELECT MAX(date_record) AS cutoff FROM historical_summary_rows"
        ).fetchone()["cutoff"]
        _flag_historically_accounted_commissions(conn, cutoff)

    return result
