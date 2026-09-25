"""
Reads a Kenjin AOR (Acknowledgment of Receipt) export and uses it to
fill in contracts' paid-date columns - the same fields Accounts used to
have to type into the Master report by hand after cross-checking this
exact file. Once those dates are on file, the existing detection
pipeline (app.commission.process_commission_run) picks up from there
exactly as if the Master report itself had carried them - nothing about
agency/agent splitting, the Date Record summary, or the review-and-
confirm workflow needed to change for this to work.

The real export's `Reference No` column is free text a staff member
typed by hand, not a structured field - see _classify_reference for
the exact rules (confirmed against a real sample together with the
business, not guessed):
  - An "(INST X/Y)" tag anywhere in the text is authoritative,
    regardless of what word precedes it - "ADVANCE BALANCE PAYMENT
    (INST 01/24)" is installment 1, not a full-payment completion.
    Only installment 1 and 6 matter for commission (the fixed release
    points - see rules.py); any other installment number is recorded
    (so the receipt is never reprocessed) but triggers nothing.
  - No INST tag at all: "FULL PAYMENT", "BALANCE PAYMENT", or "EARLY
    SETTLEMENT" all mean this receipt completes the full price (staff
    confirmed "early settlement" = the customer paid off the whole
    remaining balance ahead of schedule, not just one installment).
    Plain "DEPOSIT" or "PARTIAL PAYMENT" (no tag, not the completing
    one) and "STAMP DUTY" are recognized but non-triggering - a step
    along the way, not a commission event.
  - Anything else doesn't match a known pattern and gets flagged for a
    human to look at, rather than silently guessed at or dropped.

The real exports routinely overlap each other (the same underlying
report, exported again later with an extended date range) and a single
export can carry several sheets that are different overlapping
snapshots of the same data - `Acknowledgment Receipt No` is what makes
re-importing safe: every receipt is only ever applied once, ever (see
aor_receipts in schema.sql), the same UNIQUE-constraint-based
idempotency historical_summary_rows already uses for Date Record rows.
"""

import datetime
import os
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import PatternFill

from . import commission
from .importer import _is_positive_whole_number, _to_iso_date, _to_number

# Matches the same green already used elsewhere in this app for "fully
# paid off" (see app/report.py's _GREEN_FILL) - confirmed against a
# real annotated AOR sample, the business's own green highlight (theme
# accent6, tint 0.8) resolves to a very close shade (~DCEDD5) to this
# one, so this keeps one consistent "paid off" color across the whole
# tool rather than introducing a second, slightly different green.
_GREEN_FILL = PatternFill(start_color="C6DEB5", end_color="C6DEB5", fill_type="solid")
_YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

_AOR_REQUIRED_HEADERS = (
    "No", "Acknowledgment Receipt No", "PO No", "Reference No", "Acknowledgment Receipt Date",
)

# Captures the run of digits/slashes/commas/ampersands/whitespace right
# after "INST" - e.g. "15/24" or "22/24, 23/24 & 24/24" - stopping
# naturally at the first character that isn't part of that run (a ")"
# in every real sample seen). Deliberately scoped to just after "INST"
# rather than searching the whole Reference No for any "N/M" pattern,
# since a plain date like "07/08/2026" elsewhere in the same text would
# otherwise falsely match as its own bogus installment number.
_INST_CLUSTER_PATTERN = re.compile(r"INST\s*([\d/,&\s]+)", re.IGNORECASE)
_INST_PAIR_PATTERN = re.compile(r"(\d+)\s*/\s*\d+")

_TRIGGER_TO_DATE_COLUMN = {
    "full_payment": "full_settlement_paid_date",
    "installment_1": "first_installment_paid_date",
    "installment_6": "sixth_installment_paid_date",
}


@dataclass
class AorReviewFlag:
    acknowledgment_receipt_no: object
    po_no: object
    message: str


@dataclass
class AorImportResult:
    receipts_seen: int = 0
    # Newly recorded this upload - 0 for a receipt already imported by
    # an earlier (possibly overlapping) AOR upload, same spirit as
    # ImportResult.historical_rows_imported in app/importer.py.
    receipts_imported: int = 0
    # How many contract paid-date columns actually got filled in - can
    # be smaller than receipts_imported, since several receipts often
    # contribute to one paid-date (see _classify_reference), and a
    # column already on file is never overwritten.
    paid_dates_written: int = 0
    # Only ever nonzero when period_start/period_end were given (see
    # import_aor_report) - a receipt whose own Acknowledgment Receipt
    # Date falls outside the chosen period, or has no date at all, so
    # it's left completely unprocessed (not recorded as seen) rather
    # than applied now - it stays available for a later upload whose
    # period actually covers it.
    receipts_outside_period: int = 0
    review_flags: list = field(default_factory=list)


def _is_aor_shaped(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=30):
        values = [cell.value for cell in row]
        if all(h in values for h in _AOR_REQUIRED_HEADERS):
            return True
    return False


def _find_header_row(sheet):
    for row in sheet.iter_rows(min_row=1, max_row=30):
        values = [cell.value for cell in row]
        if all(h in values for h in _AOR_REQUIRED_HEADERS):
            return row[0].row
    return None


def _read_aor_rows(sheet):
    """Yields each data row as a dict keyed by header, stopping at the
    first row whose "No" isn't a positive whole number - the same
    "Total"/"Contra"/repeated-header-row footer shape confirmed on the
    real Master report also shows up verbatim at the bottom of a real
    AOR export."""
    header_row_num = _find_header_row(sheet)
    if header_row_num is None:
        return
    headers = [cell.value for cell in sheet[header_row_num]]
    for row in sheet.iter_rows(min_row=header_row_num + 1):
        values = [cell.value for cell in row]
        row_no = values[headers.index("No")]
        if not _is_positive_whole_number(row_no):
            break
        yield dict(zip(headers, values))


def _classify_reference(reference_text):
    """
    Returns one of:
      ("installments", [1, 6, ...])  - an (INST X/Y) tag was found;
                                        every installment number present
      ("full_payment", None)         - completes the full price
      ("skip", None)                 - recognized, non-triggering
      ("unrecognized", None)         - doesn't match any known pattern
    See the module docstring for the confirmed rules this encodes.
    """
    if not reference_text:
        return ("unrecognized", None)
    text = str(reference_text)
    upper = text.upper()

    inst_match = _INST_CLUSTER_PATTERN.search(text)
    if inst_match:
        numbers = [int(n) for n in _INST_PAIR_PATTERN.findall(inst_match.group(1))]
        if numbers:
            return ("installments", numbers)

    # "EARLY SETTLEMENT" (no INST tag) means the customer paid off the
    # whole remaining balance ahead of schedule - confirmed with the
    # business - same completing-the-full-price meaning as "FULL
    # PAYMENT"/"BALANCE PAYMENT", just different wording.
    if "FULL PAYMENT" in upper or "BALANCE PAYMENT" in upper or "EARLY SETTLEMENT" in upper:
        return ("full_payment", None)
    # "PATRIAL PAYMENT" is a real, recurring typo in the actual export
    # (confirmed against a real sample) - same meaning as "PARTIAL
    # PAYMENT", not a different, unrecognized case.
    if "STAMP DUTY" in upper or "DEPOSIT" in upper or "PARTIAL PAYMENT" in upper or "PATRIAL PAYMENT" in upper:
        return ("skip", None)
    return ("unrecognized", None)


def _targets_for_classification(kind, numbers):
    """Maps a _classify_reference result to the set of trigger_types
    (from _TRIGGER_TO_DATE_COLUMN) this receipt actually contributes
    to. An installment number other than 1 or 6 contributes to none -
    only those two are fixed commission-release points (see rules.py)
    - but the receipt is still recorded as seen either way."""
    if kind == "full_payment":
        return {"full_payment"}
    if kind == "installments":
        targets = set()
        if 1 in numbers:
            targets.add("installment_1")
        if 6 in numbers:
            targets.add("installment_6")
        return targets
    return set()


def import_aor_report(conn, file_path, imported_by_user=None, aor_upload_id=None,
                       period_start=None, period_end=None):
    """
    Reads every AOR-shaped sheet in the workbook (a real export often
    has several overlapping ones - see the module docstring), fills in
    whichever of a contract's full_settlement_paid_date /
    first_installment_paid_date / sixth_installment_paid_date columns
    are still blank, and returns an AorImportResult with counts and
    review flags.

    Deliberately never overwrites a paid-date that's already on file
    (whether it came from a Master report upload or an earlier AOR
    one) - this only ever fills in a blank, exactly like Accounts
    manually cross-checking this file and typing the date in once.

    Upload order is deliberately NOT something the caller has to get
    right: a receipt that would genuinely trigger a commission (a real
    installment 1/6 or full-payment reference) but whose PO isn't in
    the ledger yet is left completely unrecorded - not written to
    aor_receipts at all - so it stays available for a future AOR
    upload once the Master report for that PO exists, even if that
    future upload is this exact same file re-uploaded unchanged.
    Without this, aor_receipts' own UNIQUE(acknowledgment_receipt_no)
    (needed to stop a real overlapping export from double-applying a
    receipt) would also permanently swallow a receipt that never
    actually got to apply anything the first time, just because the
    AOR file happened to be uploaded before the Master report.

    aor_upload_id: the aor_uploads row this call is processing on
    behalf of (see app.pipeline.process_aor_upload), stamped onto
    every aor_receipts row this call writes purely as a record of which
    upload first introduced each receipt - never consulted to decide
    what gets imported or shown anywhere (that's already correctly
    scoped via aor_receipts' own UNIQUE(acknowledgment_receipt_no) for
    imports, and by reading the file directly for annotate_aor_file).
    Optional - None for a caller that has no aor_uploads row yet (e.g.
    a script or test calling this directly).

    period_start/period_end: ISO date strings (both required together,
    or both left None). The real Kenjin AOR export is never cut on
    clean calendar-month boundaries (one real export covered 1 June -
    17 Aug, the next 18 Aug - 23 Sep) - staff process month by month
    regardless, so a receipt whose own Acknowledgment Receipt Date
    falls outside the given period is left completely untouched: not
    recorded into aor_receipts, not applied to a paid-date, nothing.
    It stays exactly as available for a future upload whose period
    actually covers it as if this upload had never mentioned it -
    critically, this is NOT the same as "already imported, skip
    forever" (that's what aor_receipts itself already guards against);
    it's "not this period's business yet." A receipt with no
    Acknowledgment Receipt Date at all is treated the same way when a
    period is given (there's no date to check it against), rather than
    flagged as it would be with no period given - see
    AorImportResult.receipts_outside_period.

    Deliberately does not commit the transaction, matching
    import_master_report - the caller decides when to commit.
    """
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    now_iso = datetime.datetime.now().isoformat()
    source_filename = file_path.split("/")[-1].split("\\")[-1]

    aor_sheets = [sheet for sheet in workbook.worksheets if _is_aor_shaped(sheet)]
    if not aor_sheets:
        raise ValueError(
            f"No sheet in '{file_path}' has the expected AOR headers "
            f"{_AOR_REQUIRED_HEADERS} - is this really an Acknowledgment "
            f"of Receipt export?"
        )

    result = AorImportResult()
    # Excludes voided rows (see void_aor_trigger) - a receipt staff
    # voided because it turned out wrong is no longer "already
    # applied," so a future upload carrying the same acknowledgment
    # receipt no (Kenjin re-exporting the same real-world receipt,
    # corrected) can be reprocessed instead of being silently skipped.
    already_imported = {
        row["acknowledgment_receipt_no"]
        for row in conn.execute("SELECT acknowledgment_receipt_no FROM aor_receipts WHERE voided_at IS NULL")
    }
    # Whether a receipt's own classification counts as a "valid
    # payment" depends on its PO actually existing - snapshotted once
    # up front rather than queried per row.
    existing_po_nos = {row["po_no"] for row in conn.execute("SELECT po_no FROM contracts")}
    seen_this_upload = set()
    # ack_no -> {po_no, receipt_date, reference_text, payment_received,
    # trigger_type} - persisted to aor_receipts below regardless of
    # what a row classified as, recording what every receipt actually
    # was (not just the ones that mattered).
    receipt_details = {}

    # (po_no, trigger_type) -> latest Acknowledgment Receipt Date seen
    # for it. Several receipts can contribute to the same trigger (e.g.
    # a split "ADVANCE PARTIAL PAYMENT" + "ADVANCE BALANCE PAYMENT",
    # both tagged the same installment number) - the later date is the
    # one that actually completed it.
    paid_date_candidates = {}

    for sheet in aor_sheets:
        for raw_row in _read_aor_rows(sheet):
            result.receipts_seen += 1
            ack_no = raw_row.get("Acknowledgment Receipt No")
            po_no = raw_row.get("PO No")

            if not ack_no:
                result.review_flags.append(AorReviewFlag(
                    None, po_no, "Row has no Acknowledgment Receipt No - can't be safely re-imported, skipped.",
                ))
                continue
            if ack_no in already_imported or ack_no in seen_this_upload:
                continue  # already applied by this or an earlier (possibly overlapping) upload

            receipt_date = _to_iso_date(raw_row.get("Acknowledgment Receipt Date"))

            if period_start is not None:
                if receipt_date is None or not (period_start <= receipt_date <= period_end):
                    # Checked (and skipped) BEFORE seen_this_upload.add -
                    # this row must not count as "applied" so a later
                    # upload whose period actually covers it can still
                    # pick it up. See the period_start/period_end
                    # docstring above for why this is deliberately
                    # different from the already_imported skip above.
                    result.receipts_outside_period += 1
                    continue

            po_no_valid = _is_positive_whole_number(po_no)
            kind, numbers = _classify_reference(raw_row.get("Reference No"))
            targets = _targets_for_classification(kind, numbers)

            if po_no_valid and receipt_date is not None and targets and int(po_no) not in existing_po_nos:
                # This receipt would matter - a real commission trigger -
                # but its PO isn't in the ledger yet (Master report for
                # it hasn't been uploaded, or was uploaded after this
                # AOR file by mistake). Leave it completely unrecorded,
                # the same way an out-of-period receipt above is, so a
                # LATER upload - even a re-upload of this exact same
                # file, once the Master report for this PO exists -
                # still picks it up. Recording it now would mark it
                # "seen" forever via aor_receipts' own UNIQUE(
                # acknowledgment_receipt_no), permanently losing this
                # payment the moment upload order goes wrong.
                result.review_flags.append(AorReviewFlag(
                    ack_no, po_no,
                    "This PO doesn't exist in the ledger yet - upload the Master report first; "
                    "this receipt will be picked up automatically on a future AOR upload.",
                ))
                continue

            seen_this_upload.add(ack_no)
            receipt_details[ack_no] = {
                "po_no": None,
                "receipt_date": receipt_date,
                "reference_text": raw_row.get("Reference No"),
                "payment_received": _to_number(raw_row.get("Payment Received (RM)")),
                "trigger_type": None,
            }

            if not po_no_valid:
                result.review_flags.append(AorReviewFlag(
                    ack_no, po_no, "Row has no valid PO No - skipped.",
                ))
                continue
            po_no = int(po_no)
            receipt_details[ack_no]["po_no"] = po_no

            if receipt_date is None:
                result.review_flags.append(AorReviewFlag(
                    ack_no, po_no, "Row has no Acknowledgment Receipt Date - skipped.",
                ))
                continue

            if kind == "unrecognized":
                result.review_flags.append(AorReviewFlag(
                    ack_no, po_no,
                    f"Reference No {raw_row.get('Reference No')!r} doesn't match any known "
                    f"pattern - not applied, needs a human to check.",
                ))

            # targets non-empty here always means po_no is already in
            # existing_po_nos - the branch above already sent the
            # opposite case (a real trigger for a PO not yet in the
            # ledger) down its own path without ever reaching here.
            if targets:
                receipt_details[ack_no]["trigger_type"] = ",".join(sorted(targets))
            for target in targets:
                key = (po_no, target)
                if key not in paid_date_candidates or receipt_date > paid_date_candidates[key]:
                    paid_date_candidates[key] = receipt_date

    # Recorded as imported regardless of what it classified as
    # (including "skip" and "unrecognized") - once a human has had the
    # chance to see an unrecognized one flagged, re-flagging the exact
    # same receipt on every future overlapping upload adds nothing.
    #
    # ON CONFLICT rather than a plain INSERT: acknowledgment_receipt_no
    # is only ever excluded from already_imported (above) when its
    # existing row was voided, so the only way this INSERT can collide
    # with an existing row is a previously-voided one getting
    # reapplied with fresh data - reuse and overwrite that exact row,
    # resetting voided_at/voided_by_user/void_reason back to NULL,
    # rather than fail on the UNIQUE constraint.
    for ack_no in seen_this_upload:
        details = receipt_details[ack_no]
        conn.execute(
            "INSERT INTO aor_receipts "
            "(acknowledgment_receipt_no, po_no, imported_at, imported_by_user, source_filename, "
            "aor_upload_id, receipt_date, reference_text, payment_received, trigger_type) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(acknowledgment_receipt_no) DO UPDATE SET "
            "po_no = excluded.po_no, imported_at = excluded.imported_at, "
            "imported_by_user = excluded.imported_by_user, source_filename = excluded.source_filename, "
            "aor_upload_id = excluded.aor_upload_id, receipt_date = excluded.receipt_date, "
            "reference_text = excluded.reference_text, payment_received = excluded.payment_received, "
            "trigger_type = excluded.trigger_type, "
            "voided_at = NULL, voided_by_user = NULL, void_reason = NULL",
            (
                ack_no, details["po_no"], now_iso, imported_by_user, source_filename,
                aor_upload_id, details["receipt_date"], details["reference_text"],
                details["payment_received"], details["trigger_type"],
            ),
        )
    result.receipts_imported = len(seen_this_upload)

    for (po_no, target), receipt_date in paid_date_candidates.items():
        contract = conn.execute("SELECT * FROM contracts WHERE po_no = ?", (po_no,)).fetchone()
        if contract is None:
            # Defensive only - every po_no reaching paid_date_candidates
            # was already confirmed to be in existing_po_nos by the
            # per-row loop above, which sends a PO not yet in the
            # ledger down its own unrecorded, retryable path instead of
            # ever getting here. Kept in case that invariant is ever
            # broken by a future change.
            result.review_flags.append(AorReviewFlag(
                None, po_no, "This PO doesn't exist in the ledger yet - upload the Master report first.",
            ))
            continue
        date_column = _TRIGGER_TO_DATE_COLUMN[target]
        if contract[date_column] is not None:
            continue  # never overwrite an existing paid-date
        conn.execute(
            f"UPDATE contracts SET {date_column} = ? WHERE po_no = ?",
            (receipt_date, po_no),
        )
        result.paid_dates_written += 1

    return result


def void_aor_trigger(conn, po_no, trigger_type, voided_by_user, reason):
    """
    Voids one PO's whole trigger (installment_1, installment_6, or
    full_payment) - every AOR receipt that contributed to it, not just
    one - because the underlying data turned out wrong (a typo'd PO
    No, a misread Reference No, a receipt that should never have
    matched this PO at all).

    Deliberately voids the WHOLE trigger rather than a single receipt:
    a real receipt (the "split payment" pattern - ADVANCE PARTIAL +
    ADVANCE BALANCE, both tagged the same installment) is sometimes
    two receipts working together, and there's no reliable way to tell
    which one of a pair is the bad one without a human reading both.
    Voiding both and letting the correct data reapply naturally on the
    next AOR upload (the real Kenjin export is cumulative - it keeps
    re-listing old receipts, not just new ones, so nothing needs a
    special "corrected" file just for this) is simpler and safer than
    guessing which single receipt to blame.

    What voiding actually does:
      - Every un-voided aor_receipts row for this po_no whose own
        trigger_type includes this trigger gets marked voided (see the
        voided_at/voided_by_user/void_reason columns in schema.sql) -
        never deleted, and no longer counted as "already applied," so
        a future AOR upload carrying the same Acknowledgment Receipt
        No can reapply it once the data's actually right.
      - The contract's own paid-date column for this trigger is reset
        to NULL, so there's nothing left over to block the correct
        date from being written once the receipt reapplies.
      - Any commission_event already raised from this trigger is
        cleared too: a confirmed one goes through the normal
        commission.void_commission_event (kept forever, with its own
        reason, same as a human voiding one from the Review page); a
        still-pending one is simply deleted (it was never sent to
        Accounts, so there's nothing there worth keeping a permanent
        record of) - either way, contracts.*_commission_flagged gets
        cleared so the trigger is ready to be detected fresh.

    A receipt whose own trigger_type happens to cover MORE than one
    trigger (e.g. a single Reference No somehow tagged with both INST
    01 and INST 06 - rare, but the classification rules allow it) gets
    voided in full even if only one of its triggers was asked for -
    same reasoning as the split-payment case: the other trigger it
    touched will simply reapply on the next AOR upload too.

    Returns True if anything was actually voided, False if no
    un-voided receipt for this po_no/trigger_type exists (a stale
    page, a typo'd PO No, or a double-submitted form).
    """
    now_iso = datetime.datetime.now().isoformat()
    candidate_rows = conn.execute(
        "SELECT id, trigger_type FROM aor_receipts WHERE po_no = ? AND voided_at IS NULL",
        (po_no,),
    ).fetchall()
    matching_ids = [
        row["id"] for row in candidate_rows
        if row["trigger_type"] and trigger_type in row["trigger_type"].split(",")
    ]
    if not matching_ids:
        return False

    for receipt_id in matching_ids:
        conn.execute(
            "UPDATE aor_receipts SET voided_at = ?, voided_by_user = ?, void_reason = ? WHERE id = ?",
            (now_iso, voided_by_user, reason, receipt_id),
        )

    date_column = _TRIGGER_TO_DATE_COLUMN[trigger_type]
    conn.execute(f"UPDATE contracts SET {date_column} = NULL WHERE po_no = ?", (po_no,))

    event = conn.execute(
        "SELECT id, status FROM commission_events WHERE po_no = ? AND trigger_type = ? "
        "AND status IN ('pending', 'confirmed')",
        (po_no, trigger_type),
    ).fetchone()
    if event is not None and event["status"] == "confirmed":
        commission.void_commission_event(conn, event["id"], voided_by_user, reason)
    else:
        if event is not None:
            conn.execute("DELETE FROM commission_events WHERE id = ?", (event["id"],))
        commission.clear_commission_flag(conn, po_no, trigger_type)

    return True


def annotate_aor_file(file_path, output, po_period_start, po_period_end):
    """
    Writes a copy of the AOR export to `output` (a path or a file-like
    object) - every original sheet left exactly as uploaded (a genuine
    1:1 copy for cross-referencing against what was actually uploaded,
    formulas and all - see the "reads the file twice" note below),
    plus three new sheets built from the SAME raw rows, all scoped by
    each row's own Purchase Statement Date (po_period_start/
    po_period_end, ISO date strings, inclusive both ends) - a different
    axis from the row's own Acknowledgment Receipt Date (payments
    received this month routinely settle POs purchased, and thus
    statemented, in an earlier one). Read straight from the file's own
    column, never looked up from the ledger - every row already
    carries its own Purchase Statement Date regardless of whether that
    PO has made it into a Master report upload yet, so a receipt for a
    PO nobody's uploaded the Master report for is still worth seeing
    here.

      - "Payments (PO Date)": every row whose own Purchase Statement
        Date falls in that range, using the file's own raw columns
        verbatim - the real AOR export's own Customer ID/Name are
        frequently blank,
        and this is deliberately NOT enriched from the ledger - this
        sheet is purely a faithful reflection of what's actually in
        the file, the same "exact copy" spirit as the untouched
        original sheets, just narrowed down. Two extra columns this
        tool adds on top: Trigger Type and Source File.
      - "Valid Payments (PO Date)": that same subset narrowed further
        to the ones that actually matter for commission - every
        receipt belonging to a PO that reached full payment in this
        file (the whole group leading up to it, not just the
        completing row - a "PARTIAL"/"PATRIAL PAYMENT" followed by a
        "BALANCE PAYMENT" for the same PO is one completed sale, and
        both rows are included), colored green, and every receipt
        whose own (INST X/Y) tag is installment 1 or 6 specifically,
        colored yellow - confirmed against a real annotated sample,
        the exact colors and grouping the business already uses by
        hand.
      - "Boundary Payments (PO Date)": a real, confirmed Kenjin quirk -
        Kenjin sometimes statements a receipt the day after the PO was
        actually purchased, which can roll it into the next calendar
        month (a PO purchased 31 Aug can get Purchase Statement Date 1
        Sep). That receipt then silently lands in the WRONG month's
        "Valid Payments (PO Date)" - this month's report looks like
        it's missing a payment that's actually sitting one day into
        the next month's. Rather than have the system guess which
        month a boundary row "really" belongs in (this file alone
        can't tell - only the ledger's own PO Date can, and even that's
        sometimes not uploaded yet), this sheet surfaces every row
        whose Purchase Statement Date is exactly one day before
        po_period_start or one day after po_period_end AND would
        otherwise have qualified for "Valid Payments (PO Date)" (same
        green/yellow rule) - so a human can glance at it and decide
        whether it belongs to this period. Deliberately NOT
        auto-included in either sheet above, and deliberately not
        auto-resolved by looking anything up in the ledger either -
        this is a case where only a human can judge which period a
        boundary row really belongs to.

    All three new sheets are sorted by PO No ascending (20260299,
    20260300, ...) rather than the file's own row order, so a PO's
    group of rows sits together and each sheet reads top to bottom in
    order.

    Deliberately NOT scoped to "only the receipts this specific upload
    newly introduced" - every row that's actually in this file and
    matches the PO-date range shows up, full stop, even if that exact
    receipt was already recorded by an earlier overlapping upload (the
    real Kenjin export is cumulative - an "August" export re-lists
    every receipt back to whenever records began). This sheet's whole
    point is being a faithful cross-reference against the file you're
    looking at right now - hiding a row that's plainly sitting in the
    file just because some earlier upload happened to see it first
    would defeat that.

    Reads the uploaded file twice on purpose: once with data_only=True
    to classify rows and pull out values (the same read every other
    function in this module uses), and once completely untouched to
    build the original sheet(s) from - so a cell that happens to hold
    a formula in the original file is never silently flattened to its
    cached value just because this function also had to read it.
    """
    values_workbook = openpyxl.load_workbook(file_path, data_only=True)
    aor_sheets = [sheet for sheet in values_workbook.worksheets if _is_aor_shaped(sheet)]

    # First pass: which POs have a full-payment completion anywhere in
    # this file, so every receipt row for that PO (not just the
    # completing one) is pulled into "Valid Payments (PO Date)" as
    # green - regardless of the PO-date filter itself, so a full
    # payment's own earlier deposit/partial rows aren't cut off just
    # because this pass hasn't reached the PO-date check yet.
    full_payment_pos = set()
    for sheet in aor_sheets:
        for raw_row in _read_aor_rows(sheet):
            po_no = raw_row.get("PO No")
            if not _is_positive_whole_number(po_no):
                continue
            kind, _ = _classify_reference(raw_row.get("Reference No"))
            if kind == "full_payment":
                full_payment_pos.add(int(po_no))

    # One day outside each end of the chosen range - see "Boundary
    # Payments (PO Date)" in the docstring above for why exactly one
    # day: it's the confirmed real-world drift (Kenjin sometimes
    # statements a receipt the calendar day after the actual PO date).
    day_before_start = (datetime.date.fromisoformat(po_period_start) - datetime.timedelta(days=1)).isoformat()
    day_after_end = (datetime.date.fromisoformat(po_period_end) + datetime.timedelta(days=1)).isoformat()

    # Second pass: build all three new sheets' rows in one walk through
    # the file, in the order rows appear - same spirit as a human
    # scrolling through it top to bottom and filtering as they go.
    headers = None
    payment_rows = []  # list of (po_no, row_values, fill_or_None)
    boundary_rows = []  # same shape, but for the one-day-outside-the-range case
    for sheet in aor_sheets:
        header_row_num = _find_header_row(sheet)
        if header_row_num is None:
            continue
        sheet_headers = [cell.value for cell in sheet[header_row_num]]
        if headers is None:
            headers = sheet_headers
        for raw_row in _read_aor_rows(sheet):
            po_no = raw_row.get("PO No")
            statement_date = _to_iso_date(raw_row.get("Purchase Statement Date"))
            in_period = statement_date is not None and po_period_start <= statement_date <= po_period_end
            on_boundary = statement_date in (day_before_start, day_after_end)
            if not in_period and not on_boundary:
                continue

            kind, numbers = _classify_reference(raw_row.get("Reference No"))
            targets = _targets_for_classification(kind, numbers)
            fill = None
            if "installment_1" in targets or "installment_6" in targets:
                fill = _YELLOW_FILL
            elif _is_positive_whole_number(po_no) and int(po_no) in full_payment_pos:
                fill = _GREEN_FILL
            trigger_type = ",".join(sorted(targets)) if targets else None
            values = [raw_row.get(h) for h in sheet_headers] + [trigger_type, os.path.basename(file_path)]

            if in_period:
                payment_rows.append((po_no, values, fill))
            elif fill is not None:
                # Boundary Payments only ever shows rows that would
                # already have qualified for Valid Payments - see the
                # docstring above for why an unrecognized/non-triggering
                # row near the edge isn't worth surfacing here.
                boundary_rows.append((po_no, values, fill))

    # Sorted by PO No ascending (20260299, 20260300, ...) rather than
    # the file's own row order, so a PO's group of rows is easy to
    # find and everything reads in one consistent order.
    payment_rows.sort(key=lambda entry: entry[0] if _is_positive_whole_number(entry[0]) else float("inf"))
    boundary_rows.sort(key=lambda entry: entry[0] if _is_positive_whole_number(entry[0]) else float("inf"))
    valid_rows = [entry for entry in payment_rows if entry[2] is not None]

    output_workbook = openpyxl.load_workbook(file_path)  # untouched - this is what gets kept as-is
    output_headers = (headers or []) + ["Trigger Type", "Source File"]

    def _write_sheet(name, rows):
        sheet_name = name
        suffix = 2
        while sheet_name in output_workbook.sheetnames:
            sheet_name = f"{name} ({suffix})"
            suffix += 1
        sheet = output_workbook.create_sheet(sheet_name)
        for col, header in enumerate(output_headers, start=1):
            sheet.cell(row=1, column=col, value=header)
        for row_offset, (_, values, fill) in enumerate(rows, start=2):
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_offset, column=col, value=value)
                if fill is not None:
                    cell.fill = fill

    if headers is not None:
        _write_sheet("Payments (PO Date)", payment_rows)
        _write_sheet("Valid Payments (PO Date)", valid_rows)
        _write_sheet("Boundary Payments (PO Date)", boundary_rows)

    output_workbook.save(output)
